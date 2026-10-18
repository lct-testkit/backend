"""Сервис импорта каталогов — 5 фаз dop.md §4.12: загрузка (уже сделана
`files`-модулем), профилирование, маппинг, dry-run, применение (+ откат).

Батчи по 500 с чекпоинтом обрабатываются не здесь, а в `imports.tasks`
(периодический скан, тот же приём, что `workflow.tasks`/`registry.tasks`):
`apply()`/`rollback()` только переводят задание в состояние
`applying`/`rolling_back` в рамках HTTP-запроса, фактическую построчную
работу продолжает фоновая задача, читая уже провалидированные `row_data` из
`import_row_results` — так апдейт с батчами переживает падение воркера
без повторного разбора файла.
"""

from __future__ import annotations

import asyncio
import copy
import datetime as dt
import hashlib
import io
import time
import uuid
from collections import Counter
from collections.abc import Iterable
from dataclasses import dataclass
from dataclasses import field as _dc_field
from decimal import Decimal
from typing import Any

import openpyxl
import structlog
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import Select, and_, delete, func, or_, select, update
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.core.context import ActorContext, set_actor
from app.core.errors import AppError, ErrorCode, ForbiddenError, NotFoundError
from app.core.ids import uuid7
from app.core.metrics import import_duration_seconds, import_rows_total
from app.core.normalize import parse_contact_methods, split_full_name
from app.core.security import Principal
from app.core.storage import (
    download_object_bytes,
    ensure_bucket,
    inspect_object,
    upload_object_bytes,
)
from app.modules.audit.actions import AuditAction
from app.modules.audit.service import AuditService
from app.modules.catalog.license_keys import LicenseKey
from app.modules.catalog.license_matching import LicenseMatcher
from app.modules.catalog.license_values import parse_signed, parse_term
from app.modules.catalog.models import Direction, Organization, OrganizationLicense, Product, Region
from app.modules.files.models import Attachment, File, FileStatus
from app.modules.identity.models import User
from app.modules.imports.fields import (
    ADMIN_ONLY_ENTITY_TYPES,
    HANDLED_ENTITY_TYPES,
    FieldSpec,
    fields_for,
    missing_mapping_labels,
    natural_key_for,
    validate_field,
)
from app.modules.imports.handlers import HANDLERS, Batch, RowCheck, RowIssue, rollback_effects
from app.modules.imports.mapping import suggest_mapping
from app.modules.imports.models import (
    ImportEntityType,
    ImportJob,
    ImportJobStatus,
    ImportMode,
    ImportPreset,
    ImportRowResult,
    ImportRowStatus,
)
from app.modules.imports.parsing import (
    ParsedTable,
    SheetInfo,
    list_sheets,
    parse_table,
    sanitize_formula,
)
from app.modules.imports.schemas import ImportOptions
from app.modules.registry.models import EgrulEntry

logger = structlog.get_logger(__name__)

# Задание, зависшее из-за «отравленной» строки, раньше вешало весь тик воркера; теперь строка
# помечается ошибкой, а причина в отчёте — понятная, без текста исключения (в нём бывает SQL).
_GENERIC_ROW_ERROR = "Не удалось применить строку из-за внутренней ошибки (подробности в журнале)"
_EDITABLE_MAPPING_STATUSES = (
    ImportJobStatus.UPLOADED.value,
    ImportJobStatus.MAPPED.value,
    ImportJobStatus.VALIDATED.value,
)

# kind -> (модель, целевое поле в row_data, колонка поиска по значению из
# файла). Третий элемент раньше был неявно `model.code` — П3 добавляет
# `organization_name`, у которого колонка поиска не `code`, а `name`, так
# что параметризовали явно (region_code/direction_code продолжают resolve'ить
# по `.code`, поведение не изменилось).
_FK_TARGETS: dict[str, tuple[type, str, Any]] = {
    "region_code": (Region, "region_id", Region.code),
    "direction_code": (Direction, "direction_id", Direction.code),
    "organization_name": (Organization, "organization_id", Organization.name),
}
_MODEL_BY_ENTITY: dict[str, type] = {
    ImportEntityType.ORGANIZATION.value: Organization,
    ImportEntityType.PRODUCT.value: Product,
    ImportEntityType.LICENSE.value: OrganizationLicense,
}
_PENDING_STATUSES = (ImportRowStatus.OK.value, ImportRowStatus.WARN.value)


def _apply_outcome(row: ImportRowResult) -> str:
    """Итог строки после `apply_batch` для `crm_import_rows_total`: применена, пропущена или ошибка.
    Применённая строка остаётся `ok`/`warn` — отличается от ожидающей только `entity_id`."""
    if row.status == ImportRowStatus.SKIPPED.value:
        return "skipped"
    if row.status == ImportRowStatus.ERROR.value:
        return "error"
    return "applied"


def _count_rows(entity_type: str, outcomes: Iterable[str]) -> None:
    for status, count in Counter(outcomes).items():
        import_rows_total.labels(entity_type=entity_type, status=status).inc(count)


# B3-5: соответствие вида поля и кода проблемы строки для `_extract_row`/`_dry_run*`. Виды, не
# упомянутые здесь, получают общий код `field_invalid` — таблица задачи называет явно только эти
# два (остальные подсказки не предполагались автором задачи как отдельные коды).
_FIELD_ISSUE_HINTS: dict[str, tuple[str, str]] = {
    "transfer_status": ("status_unknown", "Допустимо: передано, в работе, не начато, отказ"),
    "license_term": (
        "term_invalid",
        "Укажите срок в годах (1–50) или год окончания (2000–2100)",
    ),
}


def _field_error_issue(spec: FieldSpec, message: str) -> RowIssue:
    code, hint = _FIELD_ISSUE_HINTS.get(spec.kind, ("field_invalid", None))
    return RowIssue(code=code, field=spec.target, message=message, hint=hint)


def issue_text(entry: Any) -> str:
    """`ImportRowResult.errors` хранит и старые строки, и новые объекты `RowIssue` (B3-5) — этот
    хелпер нормализует обе формы в читаемый текст."""
    if isinstance(entry, dict):
        text = str(entry.get("message", ""))
        hint = entry.get("hint")
        return f"{text} ({hint})" if hint else text
    return str(entry)


def _issue_dict(issue: RowIssue) -> dict[str, Any]:
    return {
        "code": issue.code,
        "field": issue.field,
        "message": issue.message,
        "hint": issue.hint,
        "severity": issue.severity,
    }


def entry_to_issue_dict(entry: Any) -> dict[str, Any]:
    """Нормализует одну запись `ImportRowResult.errors` (строка или объект) в форму
    `RowIssueOut` — для `ImportRowOut.issues` и построения сводки/отчёта."""
    if isinstance(entry, dict):
        return {
            "code": entry.get("code", "other"),
            "field": entry.get("field"),
            "message": entry.get("message", ""),
            "hint": entry.get("hint"),
            "severity": entry.get("severity", "error"),
        }
    return {
        "code": "other",
        "field": None,
        "message": str(entry),
        "hint": None,
        "severity": "error",
    }


def _json_safe(value: Any) -> Any:
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, dt.date | dt.datetime):
        return value.isoformat()
    if isinstance(value, uuid.UUID):
        return str(value)
    return value


@dataclass(slots=True)
class _ReportRow:
    """B3-5: одна проблемная строка для отчёта `_write_error_report` — лист и абсолютный номер
    строки Excel (не порядковый номер среди обработанных), статус и структурированные причины."""

    source_row: int
    sheet: str
    cells: list[str]
    status_label: str
    issues: list[RowIssue]


# B3-5: не больше 50 кодов и 5 примеров номеров строк на код — таблица задачи ограничивает размер
# `summary`, иначе файл с тысячами однотипных ошибок разрастит JSONB задания без пользы.
_SUMMARY_MAX_CODES = 50
_SUMMARY_MAX_EXAMPLES = 5


def _bold_row(sheet: Worksheet, labels: list[str], font: Font) -> list[WriteOnlyCell]:
    """Шапка листа жирным — в `write_only=True` книге ячейку нельзя стилизовать после `append()`,
    только передать уже оформленный `WriteOnlyCell`."""
    cells = []
    for label in labels:
        cell = WriteOnlyCell(sheet, value=label)
        cell.font = font
        cells.append(cell)
    return cells


def _would_bucket(exists: bool, mode: str) -> str:
    """Что сделает `apply()` со строкой, по режиму и тому, найдена ли уже запись — для `summary.
    would`. Строки-ошибки сюда не попадают (`apply()` их не трогает)."""
    if mode == ImportMode.INSERT.value:
        return "skip" if exists else "create"
    if mode == ImportMode.UPDATE.value:
        return "update" if exists else "skip"
    return "update" if exists else "create"  # upsert


def _accumulate_issues(
    by_issue: dict[str, dict[str, Any]], issues: list[RowIssue], source_row: int
) -> None:
    for issue in issues:
        entry = by_issue.setdefault(
            issue.code, {"message": issue.message, "rows": 0, "examples": []}
        )
        entry["rows"] += 1
        if len(entry["examples"]) < _SUMMARY_MAX_EXAMPLES:
            entry["examples"].append(source_row)


def _finalize_summary(
    by_issue: dict[str, dict[str, Any]],
    would: dict[str, dict[str, int]],
    duplicate_headers: list[str] | None = None,
) -> dict[str, Any]:
    ranked = sorted(by_issue.items(), key=lambda kv: kv[1]["rows"], reverse=True)
    summary: dict[str, Any] = {
        "by_issue": [{"code": code, **entry} for code, entry in ranked[:_SUMMARY_MAX_CODES]],
        "would": would,
    }
    if duplicate_headers:
        # B3-M: заголовок, повторённый в файле — колонка читается один раз (последняя), первая
        # теряется молча без этого предупреждения.
        summary["duplicate_headers"] = duplicate_headers
    return summary


@dataclass(slots=True)
class ProfileResult:
    headers: list[str]
    sample_rows: list[list[str]]
    suggested_mapping: dict[str, str]
    total_rows: int
    applied_preset: str | None = None
    #: B3-2: листы книги (пусто для CSV/JSON), выбранный лист/строка заголовка (из query-параметров
    #: или уже сохранённых `job.options`, если query их не переопределяет) и подсказка автопоиска.
    sheets: list[SheetInfo] = _dc_field(default_factory=list)
    sheet: str | None = None
    header_row: int = 1
    suggested_header_row: int = 1


class ImportService:
    def __init__(self, session: AsyncSession) -> None:
        self._session = session
        self._audit = AuditService(session)

    # -- служебное ------------------------------------------------------

    def list_query(self) -> Select[tuple[ImportJob]]:
        return select(ImportJob)

    async def get_or_404(self, job_id: uuid.UUID) -> ImportJob:
        job = await self._session.get(ImportJob, job_id)
        if job is None:
            raise NotFoundError("Задание импорта", job_id)
        return job

    async def _get_ready_file(self, file_id: uuid.UUID) -> File:
        file = await self._session.get(File, file_id)
        if file is None or file.deleted_at is not None:
            raise NotFoundError("Файл", file_id)
        if file.status != FileStatus.READY.value:
            raise AppError(ErrorCode.VALIDATION, "Файл ещё не прошёл проверку")
        return file

    # -- фаза 1: загрузка (создание задания на уже готовый file_id) ----

    async def create_job(
        self,
        principal: Principal,
        *,
        file_id: uuid.UUID,
        entity_type: str,
        mode: str,
        source_format: str,
    ) -> ImportJob:
        if entity_type not in _MODEL_BY_ENTITY and entity_type not in HANDLED_ENTITY_TYPES:
            raise AppError(ErrorCode.VALIDATION, f"Импорт типа {entity_type!r} не поддерживается")
        # B3-8: направления и список ответственных меняют каталог/видимость вузов напрямую (без
        # формы с привычными проверками скоупа) — только ADMIN. Право «владелец вуза» для
        # руководителя (HEAD) появится в спринте B4, здесь только самый узкий безопасный вариант.
        if entity_type in ADMIN_ONLY_ENTITY_TYPES and not principal.is_admin:
            raise ForbiddenError("Импорт этого типа доступен только администратору")
        file = await self._get_ready_file(file_id)
        # Профиль файла отдаёт его первые строки; без этой проверки руководитель мог «импортировать»
        # любой готовый файл системы (чужое вложение сделки) и читать его превью. Чужой файл для
        # него — «не найден», как и несуществующий: наличие файла не раскрывается.
        if not principal.is_admin and file.uploaded_by != principal.user_id:
            raise NotFoundError("Файл", file_id)

        job = ImportJob(
            file_id=file_id,
            entity_type=entity_type,
            mode=mode,
            source_format=source_format,
            status=ImportJobStatus.UPLOADED.value,
            initiated_by=principal.user_id,
        )
        self._session.add(job)
        await self._session.flush()
        return job

    # -- фаза 2: профилирование и подсказка маппинга --------------------

    async def _download_checked(self, job: ImportJob) -> tuple[File, bytes]:
        file = await self._get_ready_file(job.file_id)

        # Раздел 4.12/3.7: лимит размера файла импорта нигде не проверялся
        # до разбора — `download_object_bytes` тянет объект целиком в память
        # одним куском, а лимит строк (`dry_run`) в принципе не может
        # сработать раньше, чем файл уже полностью скачан и распарсен.
        # `inspect_object` — тот же приём, что уже использует
        # `files.service.commit`: читает объект потоково, не материализуя
        # его целиком, только чтобы узнать реальный размер и сходу отсеять
        # то, что заведомо превышает лимит.
        settings = get_settings()
        inspection = await inspect_object(bucket=file.bucket, key=file.storage_key)
        if inspection.size_bytes > settings.import_max_file_size_bytes:
            raise AppError(
                ErrorCode.IMPORT_BAD_FORMAT,
                "Файл превышает допустимый размер импорта: "
                f"{inspection.size_bytes} > {settings.import_max_file_size_bytes}",
            )

        content = await download_object_bytes(bucket=file.bucket, key=file.storage_key)
        return file, content

    async def _load_table(self, job: ImportJob) -> ParsedTable:
        _file, content = await self._download_checked(job)
        options = ImportOptions.model_validate(job.options or {})
        # B8-11: разбор xlsx/csv — чистый Python (openpyxl/csv), секунды CPU на крупный файл.
        # Внутри цикла событий это держало бы реплику api недоступной для всех остальных
        # запросов на всё время разбора (гипотеза B8-5, не измерено этой задачей, но правка
        # безвредна и дёшева: `to_thread` всегда корректен для CPU-синхронного кода).
        # B3-2: `sheet`/`sheets`/`header_row`/`fill_down` — выбор пользователя, сохранённый на
        # задании `PUT .../mapping`; без них разбор буквально как раньше (первый лист, первая
        # строка) — `job.options` по умолчанию `{}`, `ImportOptions()` даёт все поля `None`/пустыми.
        return await asyncio.to_thread(
            parse_table,
            content,
            source_format=job.source_format,
            sheet=options.sheet,
            sheets=options.sheets,
            header_row=options.header_row,
            fill_down=options.fill_down,
        )

    async def profile(
        self, job: ImportJob, *, sheet: str | None = None, header_row: int | None = None
    ) -> ProfileResult:
        """Раздел 4.12, фаза 2-3: превью первых 100 строк и автоподбор колонок. B3-2: `sheet`/
        `header_row` — предпросмотр ДРУГОГО листа/строки заголовка, чем уже сохранено на задании
        (если вообще сохранено); ничего здесь не пишется в БД — сохраняет только `PUT .../mapping`
        своим параметром `options`."""
        fields = fields_for(job.entity_type)
        saved_options = ImportOptions.model_validate(job.options or {})
        resolved_sheet = sheet if sheet is not None else saved_options.sheet
        resolved_header_row = header_row if header_row is not None else saved_options.header_row

        _file, content = await self._download_checked(job)
        fmt = job.source_format.lower().lstrip(".")
        sheets_info: list[SheetInfo] = []
        suggested_header_row = 1
        if fmt in ("xlsx", "xls"):
            expected_labels = [f.label for f in fields]
            sheets_info = await asyncio.to_thread(
                list_sheets, content, source_format=fmt, expected_labels=expected_labels
            )
            if sheets_info:
                if resolved_sheet:
                    match = next((s for s in sheets_info if s.name == resolved_sheet), None)
                    suggested_header_row = match.suggested_header_row if match else 1
                else:
                    visible = [s for s in sheets_info if s.visible] or sheets_info
                    suggested_header_row = visible[0].suggested_header_row

        table = await asyncio.to_thread(
            parse_table,
            content,
            source_format=job.source_format,
            sheet=resolved_sheet,
            header_row=resolved_header_row,
        )
        mapping = suggest_mapping(table.headers, list(fields), job.entity_type)

        # dop.md §4.12, фаза 3: «второй импорт того же реестра не требует
        # ручной работы» — точное совпадение заголовка с уже сохранённым
        # пресетом важнее нечёткого угадывания по словарю синонимов.
        # B3-9: только пресет АВТОРА этого задания — иначе один пользователь молча получал
        # маппинг, который сохранил кто-то другой для того же типа сущности.
        preset = await self._latest_preset(job.entity_type, created_by=job.initiated_by)
        applied_preset: str | None = None
        if preset is not None:
            allowed = {f.target for f in fields}
            for header, target in preset.mapping.items():
                if header in table.headers and target in allowed:
                    mapping[header] = target
                    applied_preset = preset.name

        return ProfileResult(
            headers=table.headers,
            sample_rows=table.rows[:100],
            suggested_mapping=mapping,
            total_rows=len(table.rows),
            applied_preset=applied_preset,
            sheets=sheets_info,
            sheet=table.sheet_name or resolved_sheet,
            header_row=table.header_row,
            suggested_header_row=suggested_header_row,
        )

    async def _latest_preset(
        self, entity_type: str, *, created_by: uuid.UUID | None = None
    ) -> ImportPreset | None:
        stmt = select(ImportPreset).where(ImportPreset.entity_type == entity_type)
        if created_by is not None:
            stmt = stmt.where(ImportPreset.created_by == created_by)
        return await self._session.scalar(stmt.order_by(ImportPreset.created_at.desc()).limit(1))

    async def get_preset_or_404(self, preset_id: uuid.UUID) -> ImportPreset:
        preset = await self._session.get(ImportPreset, preset_id)
        if preset is None:
            raise NotFoundError("Пресет маппинга", preset_id)
        return preset

    async def delete_preset(self, preset: ImportPreset) -> None:
        """B9-8 (OPEN-ISSUES B-15): раньше пресет можно было только создать (`save_mapping`,
        `save_as_preset`), убрать — только прямым SQL. Жёсткое удаление, не мягкое: пресет —
        сохранённый маппинг колонок, а не бизнес-сущность с историей переходов (как шаблон
        уведомления, `notification.service.NotificationTemplateService.delete`)."""
        preset_id = preset.id
        await self._session.delete(preset)
        await self._session.flush()
        await self._audit.record(
            AuditAction.IMPORT_PRESET_DELETED,
            entity_type="import_preset",
            entity_id=preset_id,
            changes={"name": {"old": preset.name, "new": None}},
        )

    # -- фаза 3: маппинг --------------------------------------------------

    async def save_mapping(
        self,
        job: ImportJob,
        *,
        mapping: dict[str, str],
        principal: Principal,
        save_as_preset: str | None,
        options: ImportOptions | None = None,
    ) -> ImportJob:
        # Раньше маппинг перезаписывался в любом статусе: `PUT` на применяемое или завершённое
        # задание сбрасывало его в `mapped` и стирало основу для отката.
        if job.status not in _EDITABLE_MAPPING_STATUSES:
            raise AppError(
                ErrorCode.IMPORT_NOT_APPLICABLE,
                "Маппинг можно менять только до применения импорта",
                extra={"status": job.status},
            )
        fields_by_target = {f.target: f for f in fields_for(job.entity_type)}
        unknown = [target for target in mapping.values() if target not in fields_by_target]
        if unknown:
            raise AppError(
                ErrorCode.IMPORT_MAPPING_INCOMPLETE,
                f"Неизвестные целевые поля в маппинге: {', '.join(unknown)}",
            )
        repeated = sorted({t for t in mapping.values() if list(mapping.values()).count(t) > 1})
        if repeated:
            raise AppError(
                ErrorCode.IMPORT_MAPPING_INCOMPLETE,
                "Одно поле указано для нескольких колонок: "
                + ", ".join(fields_by_target[t].label for t in repeated),
            )
        missing = missing_mapping_labels(job.entity_type, set(mapping.values()))
        if missing:
            raise AppError(
                ErrorCode.IMPORT_MAPPING_INCOMPLETE,
                "В маппинге не хватает обязательных полей: " + ", ".join(missing),
                extra={"missing": missing},
            )

        job.mapping = mapping
        # B3-2: опции листа/шапки/`fill_down` сохраняются вместе с маппингом — `_load_table`
        # перечитывает файл на каждом dry-run/apply и должна помнить выбор пользователя.
        job.options = options.model_dump(mode="json") if options is not None else {}
        job.status = ImportJobStatus.MAPPED.value
        await self._session.flush()

        if save_as_preset:
            self._session.add(
                ImportPreset(
                    name=save_as_preset,
                    entity_type=job.entity_type,
                    mapping=mapping,
                    options=job.options or None,
                    created_by=principal.user_id,
                )
            )
            await self._session.flush()
            await self._audit.record(
                AuditAction.IMPORT_PRESET_CREATED,
                entity_type="import_preset",
                entity_id=None,
                changes={"name": {"old": None, "new": save_as_preset}},
            )
        return job

    # -- фаза 4: dry-run ---------------------------------------------------

    async def dry_run(self, job: ImportJob) -> ImportJob:
        started = time.perf_counter()
        job = await self._dry_run(job)
        import_duration_seconds.labels(entity_type=job.entity_type, phase="validate").observe(
            time.perf_counter() - started
        )
        return job

    async def _dry_run(self, job: ImportJob) -> ImportJob:
        if job.status not in (ImportJobStatus.MAPPED.value, ImportJobStatus.VALIDATED.value):
            raise AppError(ErrorCode.IMPORT_MAPPING_INCOMPLETE, "Сначала сохраните маппинг колонок")

        fields_by_target = {f.target: f for f in fields_for(job.entity_type)}

        table = await self._load_table(job)
        settings = get_settings()
        if len(table.rows) > settings.import_max_rows:
            # B3-M: лимит проверяется ПОСЛЕ склейки листов (B3-2) — `table.rows` уже содержит
            # строки всех листов `sheets=["*"]`; называем листы, иначе непонятно, откуда взялось
            # число больше, чем в одном листе на экране.
            sheet_names = sorted({s for s in table.sheet_names if s})
            where = f" (листы: {', '.join(sheet_names)})" if len(sheet_names) > 1 else ""
            raise AppError(
                ErrorCode.IMPORT_BAD_FORMAT,
                f"Слишком много строк: {len(table.rows)} > {settings.import_max_rows}{where}",
            )

        # B3-M: `{h: i for i, h in enumerate(...)}` оставляет последнюю колонку с повторным
        # заголовком — первая терялась молча. Предупреждение не блокирует маппинг (та же колонка
        # просто читается один раз): пользователь узнаёт о дубле из `summary`, а не из пропавших
        # без следа данных.
        duplicate_headers = sorted({h for h in table.headers if h and table.headers.count(h) > 1})
        header_index = {h: i for i, h in enumerate(table.headers) if h}
        column_map: list[tuple[int, FieldSpec]] = [
            (header_index[header], fields_by_target[target])
            for header, target in job.mapping.items()
            if header in header_index and target in fields_by_target
        ]
        if job.entity_type in HANDLED_ENTITY_TYPES:
            return await self._dry_run_handled(job, table, column_map, duplicate_headers)

        natural_key = natural_key_for(job.entity_type)
        key_spec = fields_by_target[natural_key]
        key_col_idx = next((idx for idx, spec in column_map if spec.target == natural_key), None)
        if key_col_idx is None:
            raise AppError(
                ErrorCode.IMPORT_MAPPING_INCOMPLETE,
                f"Колонка для ключевого поля «{key_spec.label}» не найдена в файле",
            )

        fk_lookup = await self._resolve_fk_codes(column_map, table.rows)

        raw_keys = {
            row[key_col_idx].strip()
            for row in table.rows
            if key_col_idx < len(row) and row[key_col_idx].strip()
        }
        model = _MODEL_BY_ENTITY[job.entity_type]
        is_organization = job.entity_type == ImportEntityType.ORGANIZATION.value
        # Раньше — `model.inn if is_organization else model.code` (годилось
        # ровно для двух типов). `natural_key` уже посчитан выше через
        # `natural_key_for(job.entity_type)` и всегда совпадает с реальным
        # именем колонки на модели (`inn`/`code`/`contract_number`), так что
        # `getattr` — чистое обобщение без изменения поведения для
        # organization/product.
        key_column = getattr(model, natural_key)
        existing_map: dict[str, uuid.UUID] = {}
        if raw_keys and job.entity_type != ImportEntityType.LICENSE.value:
            # Мягко удалённые записи не считаются существующими: иначе строка объявляла «уже есть»
            # запись, которой пользователь не видит, а применение находило бы её же и «обновляло».
            key_stmt = select(key_column, model.id).where(
                key_column.in_(raw_keys), model.deleted_at.is_(None)
            )
            rows = (await self._session.execute(key_stmt)).all()
            existing_map = dict(rows)

        # B3-1 (BUG-9): у лицензии составной ключ (вуз, вендор, ПО, номер
        # договора) — проверка «уже существует» строится по нему, а не по
        # одному `contract_number`, иначе dry-run расходился бы с тем, что
        # реально делает `_apply_license_row` (тот же составной ключ).
        license_existing: dict[tuple[uuid.UUID, str, str, str], uuid.UUID] = {}
        if job.entity_type == ImportEntityType.LICENSE.value:
            license_existing = await self._resolve_license_existing(
                column_map, fk_lookup, table.rows
            )

        egrul_map: dict[str, EgrulEntry] = {}
        if job.entity_type == ImportEntityType.ORGANIZATION.value and raw_keys:
            entries = (
                (
                    await self._session.execute(
                        select(EgrulEntry).where(EgrulEntry.inn.in_(raw_keys))
                    )
                )
                .scalars()
                .all()
            )
            egrul_map = {e.inn: e for e in entries}

        await self._session.execute(
            delete(ImportRowResult).where(ImportRowResult.import_job_id == job.id)
        )

        seen_keys: dict[object, int] = {}
        ok = warn = error = 0
        report_rows: list[_ReportRow] = []
        by_issue: dict[str, dict[str, Any]] = {}
        would: dict[str, dict[str, int]] = {}
        is_license = job.entity_type == ImportEntityType.LICENSE.value
        source_rows = table.source_rows or list(range(1, len(table.rows) + 1))
        sheet_names = table.sheet_names or [""] * len(table.rows)

        for row_number, row in enumerate(table.rows, start=1):
            source_row = (
                source_rows[row_number - 1] if row_number <= len(source_rows) else row_number
            )
            sheet_name = sheet_names[row_number - 1] if row_number <= len(sheet_names) else ""
            row_warnings: list[str] = []
            row_issues: list[RowIssue] = []
            row_data, field_errors = self._extract_row(
                row, column_map, fk_lookup, row_warnings, row_issues
            )
            notes: list[str] = [*field_errors, *row_warnings]
            row_issues += [
                RowIssue(code="other", severity="warning", message=w) for w in row_warnings
            ]
            # B3-1: номер договора у лицензии — необязательное поле (составной ключ вуз-вендор-ПО-
            # договор, `catalog.license_keys.LicenseKey`, допускает пустой номер). «Ключевое поле»
            # для определения пустой/заполненной строки у лицензии — вуз (`organization_id`), а не
            # `natural_key` (`contract_number`, остаётся как ключ дедупликации только для
            # organization/product, у которых он действительно обязателен).
            key_value: object = (
                row_data.get("organization_id") if is_license else row_data.get(natural_key)
            )
            dedupe_key: object = key_value
            license_key_tuple: tuple[uuid.UUID, str, str, str] | None = None
            if is_license and row_data.get("organization_id"):
                license_key_tuple = LicenseKey.build(
                    uuid.UUID(row_data["organization_id"]),
                    vendor=row_data.get("vendor"),
                    product_name=row_data.get("product_name"),
                    contract_number=row_data.get("contract_number"),
                ).as_tuple()
                dedupe_key = license_key_tuple

            existing_id: uuid.UUID | None = None
            if not key_value:
                # У лицензии это не случится без своей причины (`org_not_found`/`required_empty`
                # уже даны `_extract_row`-ом выше) — отдельная ошибка не дублируется.
                if not is_license:
                    msg = f"Не удалось определить ключевое поле «{key_spec.label}»"
                    notes.append(msg)
                    row_issues.append(
                        RowIssue(
                            code="required_empty",
                            field=natural_key,
                            message=msg,
                            hint="Заполните ячейку",
                        )
                    )
                status = ImportRowStatus.ERROR.value
            else:
                status = ImportRowStatus.OK.value
                if dedupe_key in seen_keys:
                    dup_text = "Дубликат ключа внутри файла"
                    notes.append(dup_text)
                    status = ImportRowStatus.WARN.value
                    row_issues.append(
                        RowIssue(
                            code="key_duplicate_in_file",
                            message=dup_text,
                            hint=f"Строка заменит строку {seen_keys[dedupe_key]}",
                        )
                    )
                else:
                    seen_keys[dedupe_key] = source_row

                existing_id = (
                    license_existing.get(license_key_tuple)
                    if is_license
                    else existing_map.get(key_value)
                )
                if existing_id is not None and job.mode == ImportMode.INSERT.value:
                    msg = "Уже существует — будет пропущена (режим «только создание»)"
                    notes.append(msg)
                    status = ImportRowStatus.WARN.value
                    row_issues.append(
                        RowIssue(code="exists_insert_skip", severity="warning", message=msg)
                    )
                elif existing_id is None and job.mode == ImportMode.UPDATE.value:
                    msg = "Запись для обновления не найдена"
                    notes.append(msg)
                    status = ImportRowStatus.ERROR.value
                    row_issues.append(RowIssue(code="update_not_found", message=msg))

                if row_warnings and status == ImportRowStatus.OK.value:
                    status = ImportRowStatus.WARN.value
                if field_errors:
                    status = ImportRowStatus.ERROR.value

                is_organization = job.entity_type == ImportEntityType.ORGANIZATION.value
                if is_organization and key_value in egrul_map:
                    self._enrich_organization_row(row_data, egrul_map[key_value])

            if status == ImportRowStatus.OK.value:
                ok += 1
            elif status == ImportRowStatus.WARN.value:
                warn += 1
            else:
                error += 1

            if status != ImportRowStatus.ERROR.value:
                bucket = _would_bucket(existing_id is not None, job.mode)
                counts = would.setdefault(bucket, {})
                counts[job.entity_type] = counts.get(job.entity_type, 0) + 1
            _accumulate_issues(by_issue, row_issues, source_row)

            self._session.add(
                ImportRowResult(
                    import_job_id=job.id,
                    row_number=row_number,
                    source_row=source_row,
                    sheet=sheet_name or None,
                    status=status,
                    row_data=row_data,
                    errors=[_issue_dict(i) for i in row_issues],
                )
            )
            if status != ImportRowStatus.OK.value:
                status_label = (
                    "Ошибка" if status == ImportRowStatus.ERROR.value else "Предупреждение"
                )
                report_rows.append(
                    _ReportRow(
                        source_row=source_row,
                        sheet=sheet_name,
                        cells=row,
                        status_label=status_label,
                        issues=row_issues,
                    )
                )

        await self._session.flush()

        job.total_rows = len(table.rows)
        job.ok_rows, job.warn_rows, job.error_rows = ok, warn, error
        job.status = ImportJobStatus.VALIDATED.value
        job.summary = _finalize_summary(by_issue, would, duplicate_headers)
        job.result_file_id = await self._write_error_report(
            job, table.headers, report_rows, job.summary["by_issue"]
        )
        await self._session.flush()

        await self._audit.record(
            AuditAction.IMPORT_VALIDATED,
            entity_type="import_job",
            entity_id=job.id,
            changes={
                "total_rows": {"old": None, "new": job.total_rows},
                "ok_rows": {"old": None, "new": ok},
                "warn_rows": {"old": None, "new": warn},
                "error_rows": {"old": None, "new": error},
                "would_create": {"old": None, "new": would.get("create", {})},
            },
        )
        return job

    async def _resolve_fk_codes(
        self, column_map: list[tuple[int, FieldSpec]], rows: list[list[str]]
    ) -> dict[str, dict[str, uuid.UUID]]:
        codes_by_kind: dict[str, set[str]] = {kind: set() for kind in _FK_TARGETS}
        for col_idx, spec in column_map:
            if spec.kind not in _FK_TARGETS:
                continue
            for row in rows:
                if col_idx < len(row) and row[col_idx].strip():
                    codes_by_kind[spec.kind].add(row[col_idx].strip())

        resolved: dict[str, dict[str, uuid.UUID]] = {}
        for kind, (model, _field, lookup_column) in _FK_TARGETS.items():
            codes = codes_by_kind[kind]
            if not codes:
                resolved[kind] = {}
                continue
            code_stmt = select(lookup_column, model.id).where(lookup_column.in_(codes))
            deleted_at = getattr(model, "deleted_at", None)
            if deleted_at is not None:
                code_stmt = code_stmt.where(deleted_at.is_(None))
            rows_found = (await self._session.execute(code_stmt)).all()
            # `dict(rows_found)`: при неуникальном значении колонки поиска
            # (например, два вуза с совпадающим `name` — в отличие от
            # `region_code`/`direction_code`, `Organization.name` не
            # уникален) побеждает последняя строка из выборки. Это тот же
            # компромисс, что и у остального импортёра: dry-run показывает
            # результат резолва в `row_data` до применения, и явную
            # неоднозначность видно на этапе проверки, а не после apply.
            resolved[kind] = dict(rows_found)
        return resolved

    async def _resolve_license_existing(
        self,
        column_map: list[tuple[int, FieldSpec]],
        fk_lookup: dict[str, dict[str, uuid.UUID]],
        rows: list[list[str]],
    ) -> dict[tuple[uuid.UUID, str, str, str], uuid.UUID]:
        """B3-1 (BUG-9): существующие лицензии по составному ключу (вуз,
        вендор, ПО, номер договора) для строк файла — тот же ключ, что
        `_apply_license_row` использует при применении, иначе dry-run и
        apply расходились бы в том, что считают «уже существует» (ровно
        это и было причиной бага: один номер договора на несколько ПО
        молча терял строки)."""
        idx_by_target = {spec.target: idx for idx, spec in column_map}
        org_idx = idx_by_target.get("organization_name")
        vendor_idx = idx_by_target.get("vendor")
        product_idx = idx_by_target.get("product_name")
        contract_idx = idx_by_target.get("contract_number")
        org_by_name = fk_lookup.get("organization_name", {})

        def _cell(row: list[str], idx: int | None) -> str:
            return row[idx].strip() if idx is not None and idx < len(row) else ""

        keys: set[tuple[uuid.UUID, str, str, str]] = set()
        for row in rows:
            org_id = org_by_name.get(_cell(row, org_idx)) if org_idx is not None else None
            if org_id is None:
                continue
            key = LicenseKey.build(
                org_id,
                vendor=_cell(row, vendor_idx),
                product_name=_cell(row, product_idx),
                contract_number=_cell(row, contract_idx),
            )
            keys.add(key.as_tuple())
        if not keys:
            return {}

        conditions = [
            and_(
                OrganizationLicense.organization_id == org_id,
                OrganizationLicense.vendor_key == vendor_key,
                OrganizationLicense.product_key == product_key,
                OrganizationLicense.contract_key == contract_key,
            )
            for org_id, vendor_key, product_key, contract_key in keys
        ]
        stmt = select(
            OrganizationLicense.organization_id,
            OrganizationLicense.vendor_key,
            OrganizationLicense.product_key,
            OrganizationLicense.contract_key,
            OrganizationLicense.id,
        ).where(OrganizationLicense.deleted_at.is_(None), or_(*conditions))
        found = (await self._session.execute(stmt)).all()
        return {(r[0], r[1], r[2], r[3]): r[4] for r in found}

    def _extract_row(
        self,
        row: list[str],
        column_map: list[tuple[int, FieldSpec]],
        fk_lookup: dict[str, dict[str, uuid.UUID]],
        warnings: list[str] | None = None,
        issues: list[RowIssue] | None = None,
    ) -> tuple[dict[str, Any], list[str]]:
        """B3-5: `issues`, если передан, получает структурированную версию `errors` (коды из
        таблицы задачи) — список-строка `errors` не меняется, старые вызывающие продолжают
        работать как раньше."""
        row_data: dict[str, Any] = {}
        errors: list[str] = []
        full_name: tuple[str | None, str | None, str | None] | None = None
        for col_idx, spec in column_map:
            raw = row[col_idx] if col_idx < len(row) else ""
            if spec.kind in _FK_TARGETS:
                stripped = raw.strip()
                if stripped:
                    _model, target_field, _lookup_column = _FK_TARGETS[spec.kind]
                    resolved_id = fk_lookup[spec.kind].get(stripped)
                    if resolved_id is None:
                        msg = f"«{spec.label}»: значение {stripped!r} не найдено"
                        errors.append(msg)
                        if issues is not None:
                            is_org = spec.kind == "organization_name"
                            issues.append(
                                RowIssue(
                                    code="org_not_found" if is_org else "fk_not_found",
                                    field=spec.target,
                                    message=msg,
                                    hint=(
                                        "Проверьте название или добавьте колонку «ИНН вуза»"
                                        if is_org
                                        else None
                                    ),
                                )
                            )
                    else:
                        row_data[target_field] = str(resolved_id)
                elif spec.required:
                    # Раньше ни один FK-вид (region_code/direction_code) не
                    # был обязательным, поэтому пустая ячейка молча
                    # пропускалась. П3 заводит первый обязательный FK
                    # (`organization_name`) — без этой ветки пустое
                    # «Название ВУЗа» проходило бы строку в `ok` без
                    # `organization_id» и падало бы уже на `apply()`.
                    msg = f"«{spec.label}» — обязательное поле"
                    errors.append(msg)
                    if issues is not None:
                        issues.append(
                            RowIssue(
                                code="required_empty",
                                field=spec.target,
                                message=msg,
                                hint="Заполните ячейку",
                            )
                        )
                continue
            if not raw.strip() and spec.required:
                # Тот же случай, что `validate_field` обработал бы сам (пустое + required), но
                # перехваченный раньше — чтобы дать код `required_empty`, а не общий
                # `field_invalid`.
                msg = f"«{spec.label}» — обязательное поле"
                errors.append(msg)
                if issues is not None:
                    issues.append(
                        RowIssue(
                            code="required_empty",
                            field=spec.target,
                            message=msg,
                            hint="Заполните ячейку",
                        )
                    )
                continue
            value, err = validate_field(spec, raw)
            if err:
                errors.append(err)
                if issues is not None:
                    issues.append(_field_error_issue(spec, err))
                continue
            if value is None:
                continue
            if spec.kind == "person_name":
                # «ФИО» одной ячейкой раскладывается на три поля; явные колонки «Фамилия»/«Имя»/
                # «Отчество» сильнее — см. слияние ниже.
                last, first, middle = split_full_name(value)
                if not last or not first:
                    msg = f"«{spec.label}»: укажите фамилию и имя"
                    errors.append(msg)
                    if issues is not None:
                        issues.append(
                            RowIssue(code="field_invalid", field=spec.target, message=msg)
                        )
                else:
                    full_name = (last, first, middle)
                continue
            if spec.kind == "contact_methods" and warnings is not None:
                _known, unknown = parse_contact_methods(raw)
                if unknown:
                    warnings.append(f"«{spec.label}»: не распознано — {', '.join(unknown)}")
            if spec.kind in ("license_signed", "license_term") and isinstance(value, dict):
                # B3-3: эти виды разбирают одну ячейку в несколько колонок
                # модели сразу (`is_signed`+`license_signed_at`,
                # `license_term_years`/`license_valid_until`/`is_perpetual`)
                # — `value` уже словарь из `validate_field`, льём его поля
                # прямо в `row_data` вместо `spec.target`. Предупреждение
                # (дата подписания в будущем, год понят как год окончания)
                # `validate_field` не несёт (её сигнатура — только
                # значение/ошибка, тот же приём, что уже у
                # `contact_methods` выше) — получаем его повторным чистым
                # разбором исходной ячейки.
                if warnings is not None:
                    parser = parse_signed if spec.kind == "license_signed" else parse_term
                    parsed, _err = parser(raw.strip())
                    if parsed is not None and parsed.warning:
                        warnings.append(f"«{spec.label}»: {parsed.warning}")
                for key, part in value.items():
                    row_data[key] = _json_safe(part)
                continue
            row_data[spec.target] = _json_safe(value)
        if full_name is not None:
            for key, part in zip(
                ("last_name", "first_name", "middle_name"), full_name, strict=True
            ):
                if part and not row_data.get(key):
                    row_data[key] = part
        return row_data, errors

    async def _dry_run_handled(
        self,
        job: ImportJob,
        table: ParsedTable,
        column_map: list[tuple[int, FieldSpec]],
        duplicate_headers: list[str] | None = None,
    ) -> ImportJob:
        """Проверка типов, строка которых порождает несколько записей (`imports.handlers`)."""
        handler = HANDLERS[job.entity_type]
        parsed: list[tuple[dict[str, Any], list[str], list[str], list[RowIssue]]] = []
        for row in table.rows:
            warnings: list[str] = []
            field_issues: list[RowIssue] = []
            row_data, field_errors = self._extract_row(row, column_map, {}, warnings, field_issues)
            parsed.append((row_data, field_errors, warnings, field_issues))
        context = await handler.prepare(self._session, [data for data, *_ in parsed])

        await self._session.execute(
            delete(ImportRowResult).where(ImportRowResult.import_job_id == job.id)
        )
        seen_keys: set[str] = set()
        ok = warn = error = 0
        report_rows: list[_ReportRow] = []
        by_issue: dict[str, dict[str, Any]] = {}
        would: dict[str, dict[str, int]] = {}
        source_rows = table.source_rows or list(range(1, len(table.rows) + 1))
        sheet_names = table.sheet_names or [""] * len(table.rows)

        for row_number, (row, (row_data, field_errors, warnings, field_issues)) in enumerate(
            zip(table.rows, parsed, strict=True), start=1
        ):
            source_row = (
                source_rows[row_number - 1] if row_number <= len(source_rows) else row_number
            )
            sheet_name = sheet_names[row_number - 1] if row_number <= len(sheet_names) else ""
            check: RowCheck = handler.check_row(row_data, context, job.mode)
            errors = [*field_errors, *check.errors]
            notes = [*warnings, *check.warnings]
            row_issues = [*field_issues, *check.issues]
            # Предупреждения `_extract_row` (неизвестный способ связи и т.п.) — без своего кода.
            row_issues += [RowIssue(code="other", severity="warning", message=w) for w in warnings]
            # Старые обработчики (`VendorContactHandler`/`PaymentHandler`/`LearnerHandler`) пишут
            # прямо в `check.errors`/`.warnings`, минуя `add_issue` — то, что уже пришло через
            # `check.issues` (формула текста у `add_issue` та же, что у `issue_text`), не дублируем.
            known_texts = {issue_text(_issue_dict(i)) for i in check.issues}
            row_issues += [
                RowIssue(code="other", message=e) for e in check.errors if e not in known_texts
            ]
            row_issues += [
                RowIssue(code="other", severity="warning", message=w)
                for w in check.warnings
                if w not in known_texts
            ]
            if handler.dedupe_in_file and check.key:
                if check.key in seen_keys:
                    dup_text = "Повтор внутри файла — применится после предыдущей строки"
                    notes.append(dup_text)
                    row_issues.append(RowIssue(code="key_duplicate_in_file", message=dup_text))
                seen_keys.add(check.key)

            if errors:
                status = ImportRowStatus.ERROR.value
                error += 1
            elif notes:
                status = ImportRowStatus.WARN.value
                warn += 1
            else:
                status = ImportRowStatus.OK.value
                ok += 1

            if status != ImportRowStatus.ERROR.value:
                bucket = _would_bucket(check.exists, job.mode)
                counts = would.setdefault(bucket, {})
                counts[job.entity_type] = counts.get(job.entity_type, 0) + 1
            _accumulate_issues(by_issue, row_issues, source_row)

            self._session.add(
                ImportRowResult(
                    import_job_id=job.id,
                    row_number=row_number,
                    source_row=source_row,
                    sheet=sheet_name or None,
                    status=status,
                    row_data=row_data,
                    errors=[_issue_dict(i) for i in row_issues],
                )
            )
            if status != ImportRowStatus.OK.value:
                status_label = (
                    "Ошибка" if status == ImportRowStatus.ERROR.value else "Предупреждение"
                )
                report_rows.append(
                    _ReportRow(
                        source_row=source_row,
                        sheet=sheet_name,
                        cells=row,
                        status_label=status_label,
                        issues=row_issues,
                    )
                )
        await self._session.flush()

        job.total_rows = len(table.rows)
        job.processed_rows = 0
        job.ok_rows, job.warn_rows, job.error_rows = ok, warn, error
        job.status = ImportJobStatus.VALIDATED.value
        job.summary = _finalize_summary(by_issue, would, duplicate_headers)
        job.result_file_id = await self._write_error_report(
            job, table.headers, report_rows, job.summary["by_issue"]
        )
        await self._session.flush()
        await self._audit.record(
            AuditAction.IMPORT_VALIDATED,
            entity_type="import_job",
            entity_id=job.id,
            changes={
                "total_rows": {"old": None, "new": job.total_rows},
                "ok_rows": {"old": None, "new": ok},
                "warn_rows": {"old": None, "new": warn},
                "error_rows": {"old": None, "new": error},
                "would_create": {"old": None, "new": would.get("create", {})},
            },
        )
        return job

    @staticmethod
    def _enrich_organization_row(row_data: dict[str, Any], entry: EgrulEntry) -> None:
        """dop.md §11.5: «строки обогащаются из реестра» — только пустые поля,
        то, что уже есть в файле, считается более точным для этой конкретной
        строки (импортёр не должен затирать явно указанные пользователем
        данные значениями из реестра)."""
        candidates = {
            "name": entry.full_name,
            "short_name": entry.short_name,
            "kpp": entry.kpp,
            "ogrn": entry.ogrn,
            "legal_address": entry.legal_address,
        }
        for field, value in candidates.items():
            if not row_data.get(field) and value:
                row_data[field] = value
        row_data["_verified_source"] = "fns_registry"
        row_data["_registry_version_id"] = (
            str(entry.registry_version_id) if entry.registry_version_id else None
        )
        row_data["_registry_status"] = entry.status

    async def _write_error_report(
        self,
        job: ImportJob,
        headers: list[str],
        report_rows: list[_ReportRow],
        by_issue: list[dict[str, Any]],
    ) -> uuid.UUID | None:
        """B3-5: лист «Ошибки» — лист книги и АБСОЛЮТНЫЙ номер строки Excel (не порядковый среди
        обработанных), статус, «что не так» и «что сделать» по-русски, затем исходные колонки;
        лист «Сводка» — то же, что в `job.summary['by_issue']`, но в файле, который можно
        посмотреть без API. Шапка жирная, первая строка закреплена, автофильтр; формулы
        обезврежены (`sanitize_formula`, раздел 4.12)."""
        if not report_rows:
            return None

        bold = Font(bold=True)
        buffer = io.BytesIO()
        workbook = openpyxl.Workbook(write_only=True)

        errors_sheet = workbook.create_sheet("Ошибки")
        header_labels = ["Лист", "Строка в файле", "Статус", "Что не так", "Что сделать", *headers]
        errors_sheet.append(_bold_row(errors_sheet, header_labels, bold))
        for report_row in report_rows:
            what = "; ".join(i.message for i in report_row.issues) or report_row.status_label
            how = "; ".join(i.hint for i in report_row.issues if i.hint)
            errors_sheet.append(
                [
                    report_row.sheet or "",
                    report_row.source_row,
                    report_row.status_label,
                    what,
                    how,
                    *[sanitize_formula(cell) for cell in report_row.cells],
                ]
            )
        errors_sheet.freeze_panes = "A2"
        errors_sheet.auto_filter.ref = f"A1:{get_column_letter(len(header_labels))}1"
        for idx, label in enumerate(header_labels, start=1):
            errors_sheet.column_dimensions[get_column_letter(idx)].width = max(
                12, min(40, len(label) + 6)
            )

        summary_sheet = workbook.create_sheet("Сводка")
        summary_sheet.append(
            _bold_row(summary_sheet, ["Код", "Сообщение", "Строк", "Примеры"], bold)
        )
        for entry in by_issue:
            summary_sheet.append(
                [
                    entry["code"],
                    entry["message"],
                    entry["rows"],
                    ", ".join(str(n) for n in entry["examples"]),
                ]
            )
        for idx, width in enumerate((16, 48, 8, 24), start=1):
            summary_sheet.column_dimensions[get_column_letter(idx)].width = width

        workbook.save(buffer)
        content = buffer.getvalue()

        settings = get_settings()
        await ensure_bucket(settings.s3_bucket_files)
        file_id = uuid7()
        storage_key = f"{file_id}/import-errors.xlsx"
        await upload_object_bytes(
            bucket=settings.s3_bucket_files,
            key=storage_key,
            body=content,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        report_file = File(
            id=file_id,
            storage_key=storage_key,
            bucket=settings.s3_bucket_files,
            original_filename="import-errors.xlsx",
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            size_bytes=len(content),
            sha256=hashlib.sha256(content).hexdigest(),
            status=FileStatus.READY.value,
            uploaded_by=job.initiated_by,
        )
        self._session.add(report_file)
        # Отдельный flush перед вложением: `File`/`Attachment` не связаны
        # ORM `relationship()` (полиморфная связь по `entity_type`/
        # `entity_id`, см. docstring `files.models`), поэтому unit-of-work
        # не всегда выводит порядок вставки из одних только `ForeignKey` —
        # без явного flush INSERT `attachments` иногда уходит раньше
        # ещё не вставленного `files` и падает по FK.
        await self._session.flush()
        self._session.add(
            Attachment(
                file_id=file_id,
                entity_type="import_job",
                entity_id=job.id,
                category="report",
                uploaded_by=job.initiated_by,
            )
        )
        await self._session.flush()
        return file_id

    # -- фаза 5: применение -------------------------------------------------

    async def start_apply(self, job: ImportJob) -> ImportJob:
        if job.status != ImportJobStatus.VALIDATED.value:
            raise AppError(
                ErrorCode.IMPORT_NOT_APPLICABLE,
                "Сначала выполните проверку (dry-run) без ошибок в маппинге",
            )
        if job.ok_rows + job.warn_rows == 0:
            raise AppError(ErrorCode.IMPORT_NOT_APPLICABLE, "Нет строк, пригодных к применению")

        job.status = ImportJobStatus.APPLYING.value
        job.started_at = dt.datetime.now(dt.UTC)
        await self._session.flush()
        await self._audit.record(
            AuditAction.IMPORT_STARTED,
            entity_type="import_job",
            entity_id=job.id,
            changes={"mode": {"old": None, "new": job.mode}},
        )

        # B3-12: небольшое задание (по умолчанию ≤500 строк — один батч) применяется сразу в этом
        # запросе, а не ждёт минутный тик `imports.tasks.sweep_import_jobs` или арк-воркера.
        # Отдельная блокировка строки не нужна: до коммита этой транзакции (конец HTTP-запроса,
        # `DbSession`) ни cron, ни арк-задача не увидят статус `applying` — гонки с ними нет.
        # Партии всё равно идут через `apply_batch` (построчный SAVEPOINT на отравленную строку,
        # как и у фоновой задачи).
        settings = get_settings()
        inline_limit = settings.import_inline_max_rows
        if inline_limit and job.total_rows <= inline_limit:
            while await self.apply_batch(job, batch_size=settings.import_batch_size):
                pass
            await self.finalize_apply_if_done(job)
        return job

    async def apply_batch(self, job: ImportJob, *, batch_size: int) -> int:
        """Обрабатывает одну партию — вызывается `imports.tasks.sweep_import_jobs`.
        Возвращает число обработанных строк (0 — партия исчерпана)."""
        pending = (
            (
                await self._session.execute(
                    select(ImportRowResult)
                    .where(
                        ImportRowResult.import_job_id == job.id,
                        ImportRowResult.status.in_(_PENDING_STATUSES),
                        ImportRowResult.entity_id.is_(None),
                    )
                    .order_by(ImportRowResult.row_number)
                    .limit(batch_size)
                )
            )
            .scalars()
            .all()
        )
        if not pending:
            return 0

        await self._set_actor(job)
        handler = HANDLERS.get(job.entity_type)
        batch = Batch(self._session, job) if handler is not None else None
        # B3-2: один `LicenseMatcher` на партию, не на строку — кандидаты (вендоры/ПО) кэшируются
        # внутри инстанса (докстринг класса), повторный `select()` на каждой из `batch_size`
        # (по умолчанию 500) строк был бы лишней нагрузкой на БД без выигрыша в точности: каталог
        # не меняется в процессе применения одной партии. `SAVEPOINT` ниже откатывает только
        # ЗАПИСИ строки-неудачи, кэш матчера — не часть этой транзакции (только читал, до начала
        # партии), откат соседней строки его не портит.
        license_matcher = (
            LicenseMatcher(self._session)
            if job.entity_type == ImportEntityType.LICENSE.value
            else None
        )
        for row in pending:
            previous_status, previous_errors = row.status, list(row.errors or [])
            try:
                # Каждая строка — в своём SAVEPOINT: «отравленная» строка (несмапленное поле,
                # слишком длинное значение, нарушение уникальности) откатывает только себя.
                # Раньше исключение обрывало партию, строка оставалась в очереди, и задание
                # зависало навсегда.
                async with self._session.begin_nested():
                    if handler is not None and batch is not None:
                        self._record_outcome(row, await handler.apply_row(batch, row))
                    elif job.entity_type == ImportEntityType.ORGANIZATION.value:
                        await self._apply_organization_row(job, row)
                    elif job.entity_type == ImportEntityType.PRODUCT.value:
                        await self._apply_product_row(job, row)
                    elif license_matcher is not None:
                        await self._apply_license_row(job, row, license_matcher)
            except Exception as exc:  # noqa: BLE001 — любая ошибка строки -> ошибка строки, не задания
                if batch is not None:
                    batch.invalidate()
                self._fail_row(job, row, exc, previous_status, previous_errors)
        job.processed_rows += len(pending)
        await self._session.flush()
        # После flush, а не в цикле: строки ушедшей в откат партии не должны попадать в счётчик.
        _count_rows(job.entity_type, (_apply_outcome(row) for row in pending))
        return len(pending)

    async def _set_actor(self, job: ImportJob) -> None:
        """Записи аудита, которые пишет фоновая задача, принадлежат инициатору импорта, а не
        «никому»: иначе в журнале не видно, кто создал сотни контактов."""
        if job.initiated_by is None:
            return
        user = await self._session.get(User, job.initiated_by)
        set_actor(ActorContext(user_id=job.initiated_by, role=user.role if user else None))

    @staticmethod
    def _record_outcome(row: ImportRowResult, outcome: Any) -> None:
        row.effects = outcome.effects
        if outcome.skipped:
            row.status = ImportRowStatus.SKIPPED.value
            row.errors = [*(row.errors or []), *outcome.notes]
            return
        row.entity_id = outcome.entity_id
        row.before_snapshot = None
        if outcome.notes:
            row.errors = [*(row.errors or []), *outcome.notes]

    @staticmethod
    def _row_error_text(exc: Exception) -> str:
        """Причина для отчёта: текст прикладной ошибки показываем, чужое исключение — нет (в нём
        бывает SQL и значения из чужих строк)."""
        if isinstance(exc, AppError):
            reasons = [f"{e.field}: {e.reason}" for e in exc.errors]
            return "; ".join([exc.detail, *reasons]) if reasons else exc.detail
        return _GENERIC_ROW_ERROR

    def _fail_row(
        self,
        job: ImportJob,
        row: ImportRowResult,
        exc: Exception,
        previous_status: str,
        previous_errors: list[Any],
    ) -> None:
        if not isinstance(exc, AppError):
            logger.exception(
                "import_row_failed",
                job_id=str(job.id),
                row_number=row.row_number,
                error=str(exc)[:200],
            )
        row.status = ImportRowStatus.ERROR.value
        row.errors = [*previous_errors, self._row_error_text(exc)]
        row.entity_id = None
        job.error_rows += 1
        if previous_status == ImportRowStatus.WARN.value:
            job.warn_rows = max(0, job.warn_rows - 1)
        else:
            job.ok_rows = max(0, job.ok_rows - 1)

    async def _apply_organization_row(self, job: ImportJob, row: ImportRowResult) -> None:
        data = dict(row.row_data)
        internal = {k: data.pop(k) for k in list(data) if k.startswith("_")}
        inn = data.get("inn")
        existing = (
            await self._session.scalar(
                select(Organization).where(
                    Organization.inn == inn, Organization.deleted_at.is_(None)
                )
            )
            if inn
            else None
        )

        if existing is None:
            if job.mode == ImportMode.UPDATE.value:
                # Не должно случиться: dry-run уже помечает такие строки
                # `error`, значит они не попадают в выборку `apply_batch`.
                # Терминальный статус на всякий случай — иначе пустая
                # строка с `entity_id IS NULL` обрабатывалась бы вечно.
                row.status = ImportRowStatus.SKIPPED.value
                row.errors = [*row.errors, "Запись для обновления не найдена"]
                return
            org = Organization(
                name=data.get("name") or f"Организация (ИНН {inn})",
                short_name=data.get("short_name"),
                org_type=data.get("org_type", "university"),
                inn=inn,
                kpp=data.get("kpp"),
                ogrn=data.get("ogrn"),
                legal_address=data.get("legal_address"),
                actual_address=data.get("actual_address"),
                region_id=uuid.UUID(data["region_id"]) if data.get("region_id") else None,
                website=data.get("website"),
                main_phone=data.get("main_phone"),
                main_email=data.get("main_email"),
                students_count=data.get("students_count"),
                source="import",
                import_job_id=job.id,
                created_by=job.initiated_by,
                owner_id=job.initiated_by,
            )
            if internal.get("_verified_source"):
                org.verified_source = internal["_verified_source"]
                org.verified_at = dt.datetime.now(dt.UTC)
                org.registry_version_id = (
                    uuid.UUID(internal["_registry_version_id"])
                    if internal.get("_registry_version_id")
                    else None
                )
                org.registry_status = internal.get("_registry_status")
            self._session.add(org)
            await self._session.flush()
            row.entity_id = org.id
            row.before_snapshot = None
            return

        if job.mode == ImportMode.INSERT.value:
            row.status = ImportRowStatus.SKIPPED.value
            row.errors = [*row.errors, "Пропущена при применении: уже существует"]
            return

        before: dict[str, Any] = {}
        for field in (
            "name",
            "short_name",
            "kpp",
            "ogrn",
            "legal_address",
            "actual_address",
            "website",
            "main_phone",
            "main_email",
            "students_count",
            "region_id",
        ):
            if field not in data:
                continue
            new_value = uuid.UUID(data[field]) if field == "region_id" else data[field]
            old_value = getattr(existing, field)
            if _json_safe(old_value) != _json_safe(new_value):
                before[field] = _json_safe(old_value)
                setattr(existing, field, new_value)
        existing.version += 1
        await self._session.flush()
        row.entity_id = existing.id
        row.before_snapshot = before or {}

    async def _apply_product_row(self, job: ImportJob, row: ImportRowResult) -> None:
        data = dict(row.row_data)
        code = data.get("code")
        existing = (
            await self._session.scalar(
                select(Product).where(Product.code == code, Product.deleted_at.is_(None))
            )
            if code
            else None
        )

        if existing is None:
            if job.mode == ImportMode.UPDATE.value:
                row.status = ImportRowStatus.SKIPPED.value
                row.errors = [*row.errors, "Запись для обновления не найдена"]
                return
            if code and await self._session.scalar(
                select(Product.id).where(Product.code == code, Product.deleted_at.is_not(None))
            ):
                # Уникальный индекс `uq_products_code` не смотрит на `deleted_at`: повторный код
                # удалённого продукта раньше давал 500 на вставке.
                raise AppError(
                    ErrorCode.DUPLICATE,
                    f"Код {code!r} занят удалённым продуктом — восстановите его или смените код",
                )
            raw_price = data.get("base_price")
            product = Product(
                code=code,
                name=data.get("name") or code,
                description=data.get("description"),
                direction_id=uuid.UUID(data["direction_id"]) if data.get("direction_id") else None,
                duration_hours=data.get("duration_hours"),
                format=data.get("format"),
                kind=data.get("kind") or "course",
                base_price=Decimal(str(raw_price)) if raw_price is not None else None,
                currency=data.get("currency") or "RUB",
                import_job_id=job.id,
            )
            self._session.add(product)
            await self._session.flush()
            row.entity_id = product.id
            row.before_snapshot = None
            return

        if job.mode == ImportMode.INSERT.value:
            row.status = ImportRowStatus.SKIPPED.value
            row.errors = [*row.errors, "Пропущена при применении: уже существует"]
            return

        before: dict[str, Any] = {}
        for field in (
            "name",
            "description",
            "direction_id",
            "duration_hours",
            "format",
            "kind",
            "base_price",
            "currency",
        ):
            if field not in data:
                continue
            if field == "direction_id":
                new_value: Any = uuid.UUID(data[field])
            elif field == "base_price":
                new_value = Decimal(str(data[field]))
            else:
                new_value = data[field]
            old_value = getattr(existing, field)
            if _json_safe(old_value) != _json_safe(new_value):
                before[field] = _json_safe(old_value)
                setattr(existing, field, new_value)
        existing.version += 1
        await self._session.flush()
        row.entity_id = existing.id
        row.before_snapshot = before or {}

    def _license_term_fields(self, data: dict[str, Any]) -> dict[str, Any]:
        """B3-3, п.6 задачи: если в строке есть `license_term_years` и
        `license_signed_at`, а `license_valid_until` не задан явно (год
        окончания/диапазон/дата уже дали бы его в `data`), досчитать
        окончание как `signed_at + N лет`. Полноценный `LicenseHandler`
        (B3-4) считал бы это на этапе проверки; здесь — та же формула
        прямо при применении, пока лицензии остаются на обычном
        natural-key пути импорта (composite-key обработчик — отдельная
        задача, см. отчёт по эпику)."""
        result: dict[str, Any] = {}
        if (
            data.get("license_valid_until") is None
            and data.get("license_term_years") is not None
            and data.get("license_signed_at")
        ):
            signed_at = dt.date.fromisoformat(data["license_signed_at"])
            try:
                result["license_valid_until"] = signed_at.replace(
                    year=signed_at.year + int(data["license_term_years"])
                )
            except ValueError:
                # 29 февраля + N лет на невисокосный год — сознательно
                # падаем на последний день февраля, а не роняем применение.
                result["license_valid_until"] = signed_at.replace(
                    month=2, day=28, year=signed_at.year + int(data["license_term_years"])
                )
        return result

    async def _apply_license_row(
        self, job: ImportJob, row: ImportRowResult, matcher: LicenseMatcher
    ) -> None:
        """B3-1/B3-3/B3-2: ключ строки — составной (вуз, вендор, ПО, номер
        договора; `catalog.license_keys.LicenseKey`), а не только
        `contract_number` (BUG-9: один договор на несколько ПО раньше
        молча терял строки). Вуз (`organization_id`, уже разрешён раньше в конвейере, к этому
        месту приходит готовым UUID) по-прежнему ищется точным совпадением — поле есть в
        маппинге наравне с остальными `_FK_TARGETS`, это не относится к B3-4. Вендор/ПО теперь
        ищутся по каталогу тем же `LicenseMatcher`, что и API лицензий (B3-6,
        `catalog.service.OrganizationLicenseService._resolve_vendor`/`_resolve_product`) —
        раньше (B3-4, до этой задачи) `vendor_id`/`product_id` оставались `NULL` для каждой
        лицензии, заведённой импортом, независимо от того, насколько однозначно текст совпадал
        с каталогом. Неоднозначное/не найденное совпадение — не ошибка строки: `vendor`/
        `product_name` (текст) сохраняются всегда, `vendor_id`/`product_id` остаются `NULL`,
        так же, как API лицензий делает при ручном вводе текстом без выбора из каталога."""
        data = dict(row.row_data)
        vendor = data.get("vendor")
        product_name = data.get("product_name")
        contract_number = data.get("contract_number")
        organization_id = uuid.UUID(data["organization_id"])
        key = LicenseKey.build(
            organization_id,
            vendor=vendor,
            product_name=product_name,
            contract_number=contract_number,
        )
        vendor_match = await matcher.match_vendor(vendor)
        vendor_id = (
            vendor_match.entity.id if vendor_match.is_unique and vendor_match.entity else None
        )
        product_match = await matcher.match_software(product_name, vendor_id=vendor_id)
        product_id = (
            product_match.entity.id if product_match.is_unique and product_match.entity else None
        )
        existing = await self._session.scalar(
            select(OrganizationLicense).where(
                OrganizationLicense.organization_id == organization_id,
                OrganizationLicense.vendor_key == key.vendor_key,
                OrganizationLicense.product_key == key.product_key,
                OrganizationLicense.contract_key == key.contract_key,
                OrganizationLicense.deleted_at.is_(None),
            )
        )

        signed_at = (
            dt.date.fromisoformat(data["license_signed_at"])
            if data.get("license_signed_at")
            else None
        )
        valid_until_raw = data.get("license_valid_until")
        term_fields = self._license_term_fields(data)
        valid_until = (
            dt.date.fromisoformat(valid_until_raw)
            if valid_until_raw
            else term_fields.get("license_valid_until")
        )

        if existing is None:
            if job.mode == ImportMode.UPDATE.value:
                row.status = ImportRowStatus.SKIPPED.value
                row.errors = [*row.errors, "Запись для обновления не найдена"]
                return
            license_ = OrganizationLicense(
                organization_id=organization_id,
                vendor=vendor,
                vendor_id=vendor_id,
                product_name=product_name,
                product_id=product_id,
                contract_number=contract_number,
                vendor_key=key.vendor_key,
                product_key=key.product_key,
                contract_key=key.contract_key,
                license_signed_at=signed_at,
                license_valid_year=data.get("license_valid_year"),
                transfer_status=data.get("transfer_status"),
                manager_full_name=data.get("manager_full_name"),
                responsible_contacts=data.get("responsible_contacts"),
                comment=data.get("comment"),
                is_signed=data.get("is_signed"),
                license_term_years=data.get("license_term_years"),
                license_valid_until=valid_until,
                is_perpetual=bool(data.get("is_perpetual", False)),
                import_job_id=job.id,
            )
            self._session.add(license_)
            await self._session.flush()
            row.entity_id = license_.id
            row.before_snapshot = None
            return

        if job.mode == ImportMode.INSERT.value:
            row.status = ImportRowStatus.SKIPPED.value
            row.errors = [*row.errors, "Пропущена при применении: уже существует"]
            return

        before: dict[str, Any] = {}
        # Поля, которые считаются всегда (не только когда они пришли в
        # `row_data`): составной ключ пересчитывается на каждое применение
        # строки, иначе переименование вендора/ПО обновило бы текст, но
        # оставило бы старый `*_key` — `GET.../organization-licenses` (B3-6)
        # и повторный импорт того же файла перестали бы находить запись.
        # B3-2: `vendor_id`/`product_id` — туда же: запись каталога могла появиться (или
        # исчезнуть) между двумя прогонами импорта того же файла, а текст вендора/ПО при этом
        # мог не измениться вовсе — опираться на наличие ключа в `data` (как обычные поля) было
        # бы неверно, пересчёт нужен при каждом применении.
        always_fields = {
            "organization_id",
            "license_signed_at",
            "license_valid_until",
            "vendor_key",
            "product_key",
            "contract_key",
            "vendor_id",
            "product_id",
        }
        new_values: dict[str, Any] = {
            "organization_id": organization_id,
            "vendor": data.get("vendor"),
            "vendor_id": vendor_id,
            "product_name": data.get("product_name"),
            "product_id": product_id,
            "license_signed_at": signed_at,
            "license_valid_year": data.get("license_valid_year"),
            "transfer_status": data.get("transfer_status"),
            "manager_full_name": data.get("manager_full_name"),
            "responsible_contacts": data.get("responsible_contacts"),
            "comment": data.get("comment"),
            "vendor_key": key.vendor_key,
            "product_key": key.product_key,
            "contract_key": key.contract_key,
        }
        if "is_signed" in data:
            new_values["is_signed"] = data["is_signed"]
        if "license_term_years" in data:
            new_values["license_term_years"] = data["license_term_years"]
        if valid_until is not None or "license_valid_until" in data:
            new_values["license_valid_until"] = valid_until
        if "is_perpetual" in data:
            new_values["is_perpetual"] = bool(data["is_perpetual"])
        for field, new_value in new_values.items():
            if field not in data and field not in always_fields:
                continue
            old_value = getattr(existing, field)
            if _json_safe(old_value) != _json_safe(new_value):
                before[field] = _json_safe(old_value)
                setattr(existing, field, new_value)
        existing.version += 1
        await self._session.flush()
        row.entity_id = existing.id
        row.before_snapshot = before or {}

    async def _tally_applied(self, job: ImportJob) -> dict[str, dict[str, int]]:
        """B3-5: факт после применения — сколько строк создали, обновили и пропустили. Для
        типов без обработчика `before_snapshot is None` отличает создание от обновления
        (`_apply_organization_row` и соседи пишут `None` только при создании); для типов с
        обработчиком `before_snapshot` всегда `None` (`_record_outcome`), поэтому там смотрим на
        `op` в `effects` — эффект с тем же `id`, что и `entity_id` строки."""
        rows = (
            await self._session.execute(
                select(
                    ImportRowResult.status,
                    ImportRowResult.entity_id,
                    ImportRowResult.before_snapshot,
                    ImportRowResult.effects,
                ).where(ImportRowResult.import_job_id == job.id)
            )
        ).all()
        created = updated = skipped = 0
        for status, entity_id, before_snapshot, effects in rows:
            if status == ImportRowStatus.SKIPPED.value:
                skipped += 1
                continue
            if entity_id is None:
                continue
            own_op = next(
                (e.get("op") for e in effects or [] if e.get("id") == str(entity_id)), None
            )
            is_update = before_snapshot is not None or own_op == "update"
            if is_update:
                updated += 1
            else:
                created += 1
        key = job.entity_type
        return {"create": {key: created}, "update": {key: updated}, "skip": {key: skipped}}

    async def finalize_apply_if_done(self, job: ImportJob) -> bool:
        """Возвращает `True`, если задание больше не в статусе `applying`."""
        remaining = await self._session.scalar(
            select(func.count(ImportRowResult.id)).where(
                ImportRowResult.import_job_id == job.id,
                ImportRowResult.status.in_(_PENDING_STATUSES),
                ImportRowResult.entity_id.is_(None),
            )
        )
        if remaining:
            return False

        job.status = (
            ImportJobStatus.COMPLETED_WITH_ERRORS.value
            if job.error_rows
            else ImportJobStatus.COMPLETED.value
        )
        job.finished_at = dt.datetime.now(dt.UTC)
        job.rollback_available = True
        # B3-5: `summary.applied` — факт после применения, рядом с `would` (что было ПРЕДСКАЗАНО
        # на dry-run); `job.summary` может быть `None` у задания, провалидированного до B3-5.
        job.summary = {**(job.summary or {}), "applied": await self._tally_applied(job)}
        await self._session.flush()
        if job.started_at is not None:
            import_duration_seconds.labels(entity_type=job.entity_type, phase="apply").observe(
                (job.finished_at - job.started_at).total_seconds()
            )
        await self._audit.record(
            AuditAction.IMPORT_APPLIED,
            entity_type="import_job",
            entity_id=job.id,
            changes={"status": {"old": None, "new": job.status}},
        )
        return True

    # -- откат ---------------------------------------------------------

    async def start_rollback(self, job: ImportJob) -> ImportJob:
        if not job.rollback_available or job.status not in (
            ImportJobStatus.COMPLETED.value,
            ImportJobStatus.COMPLETED_WITH_ERRORS.value,
        ):
            raise AppError(ErrorCode.IMPORT_NOT_ROLLBACKABLE, "Импорт нельзя откатить")
        job.status = ImportJobStatus.ROLLING_BACK.value
        job.rollback_available = False
        # Строки, заблокированные прошлой попыткой отката, берутся снова: блокер (сделка, связанная
        # с созданным контактом) мог быть снят, и повторный откат должен добрать остальное.
        await self._session.execute(
            update(ImportRowResult)
            .where(
                ImportRowResult.import_job_id == job.id,
                ImportRowResult.status == ImportRowStatus.ROLLBACK_BLOCKED.value,
            )
            .values(status=ImportRowStatus.OK.value)
        )
        await self._session.flush()
        return job

    async def rollback_batch(self, job: ImportJob, *, batch_size: int) -> int:
        pending = (
            (
                await self._session.execute(
                    select(ImportRowResult)
                    .where(
                        ImportRowResult.import_job_id == job.id,
                        ImportRowResult.status.in_(_PENDING_STATUSES),
                        ImportRowResult.entity_id.is_not(None),
                    )
                    # С конца файла к началу: позднюю строку, которая пользовалась записью ранней
                    # (второй заказ того же человека), снимаем раньше — иначе ранняя увидит её как
                    # «зависимую» и откат заблокируется зря.
                    .order_by(ImportRowResult.row_number.desc())
                    .limit(batch_size)
                )
            )
            .scalars()
            .all()
        )
        if not pending:
            return 0

        await self._set_actor(job)
        for row in pending:
            previous_errors = list(row.errors or [])
            try:
                async with self._session.begin_nested():
                    await self._rollback_row(job, row)
            except Exception as exc:  # noqa: BLE001 — строка не откатилась, остальные идут дальше
                logger.exception(
                    "import_row_rollback_failed",
                    job_id=str(job.id),
                    row_number=row.row_number,
                    error=str(exc)[:200],
                )
                row.status = ImportRowStatus.ROLLBACK_BLOCKED.value
                row.errors = [*previous_errors, "Откат строки не удался: внутренняя ошибка"]
        await self._session.flush()
        _count_rows(job.entity_type, (row.status for row in pending))
        return len(pending)

    async def _rollback_row(self, job: ImportJob, row: ImportRowResult) -> None:
        if job.entity_type in HANDLED_ENTITY_TYPES:
            effects = copy.deepcopy(row.effects or [])
            blocked = await rollback_effects(self._session, effects)
            row.effects = effects  # новый объект: JSONB не отслеживает правку на месте
            if blocked:
                row.status = ImportRowStatus.ROLLBACK_BLOCKED.value
                row.errors = [*(row.errors or []), *blocked]
            else:
                row.status = ImportRowStatus.ROLLED_BACK.value
            return

        model = _MODEL_BY_ENTITY[job.entity_type]
        entity = await self._session.get(model, row.entity_id)
        if entity is None:
            row.status = ImportRowStatus.ROLLED_BACK.value
            return
        # `None` — запись создана этим импортом; `{}` — импорт запись обновил без изменений
        # (upsert того же значения). Раньше оба случая проверялись как `if before_snapshot`, и откат
        # «пустого» обновления удалял запись, созданную предыдущим импортом.
        if row.before_snapshot is not None:
            for field, old_value in row.before_snapshot.items():
                if (
                    field in ("region_id", "direction_id", "organization_id")
                    and old_value is not None
                ):
                    old_value = uuid.UUID(old_value)
                if field == "base_price" and old_value is not None:
                    old_value = Decimal(str(old_value))
                if field in ("license_signed_at", "license_valid_until") and old_value is not None:
                    old_value = dt.date.fromisoformat(old_value)
                setattr(entity, field, old_value)
            if row.before_snapshot:
                entity.version += 1
            row.status = ImportRowStatus.ROLLED_BACK.value
            return

        if await self._has_critical_dependents(job.entity_type, entity.id):
            row.status = ImportRowStatus.ROLLBACK_BLOCKED.value
            row.errors = [*(row.errors or []), "Откат заблокирован: есть связанные сделки"]
        else:
            entity.deleted_at = dt.datetime.now(dt.UTC)
            row.status = ImportRowStatus.ROLLED_BACK.value

    async def _has_critical_dependents(self, entity_type: str, entity_id: uuid.UUID) -> bool:
        if entity_type == ImportEntityType.ORGANIZATION.value:
            from app.modules.crm.models import Deal

            return bool(
                await self._session.scalar(
                    select(Deal.id)
                    .where(Deal.organization_id == entity_id, Deal.deleted_at.is_(None))
                    .limit(1)
                )
            )
        if entity_type == ImportEntityType.PRODUCT.value:
            from app.modules.crm.models import DealProduct

            return bool(
                await self._session.scalar(
                    select(DealProduct.id).where(DealProduct.product_id == entity_id).limit(1)
                )
            )
        return False

    async def finalize_rollback_if_done(self, job: ImportJob) -> bool:
        remaining = await self._session.scalar(
            select(func.count(ImportRowResult.id)).where(
                ImportRowResult.import_job_id == job.id,
                ImportRowResult.status.in_(_PENDING_STATUSES),
                ImportRowResult.entity_id.is_not(None),
            )
        )
        if remaining:
            return False

        blocked = await self._session.scalar(
            select(func.count(ImportRowResult.id)).where(
                ImportRowResult.import_job_id == job.id,
                ImportRowResult.status == ImportRowStatus.ROLLBACK_BLOCKED.value,
            )
        )
        if blocked:
            # Откат «успешен» только когда откатилось всё. Пока есть заблокированные строки,
            # задание остаётся завершённым с ошибками, а откат доступен снова — после разбора
            # блокеров (сделки по созданному контакту, продукт в чужой сделке).
            job.status = ImportJobStatus.COMPLETED_WITH_ERRORS.value
            job.rollback_available = True
            await self._session.flush()
            await self._audit.record(
                AuditAction.IMPORT_ROLLBACK,
                entity_type="import_job",
                entity_id=job.id,
                changes={"rollback_blocked_rows": {"old": None, "new": int(blocked)}},
            )
            return True

        job.status = ImportJobStatus.ROLLED_BACK.value
        job.rolled_back_at = dt.datetime.now(dt.UTC)
        await self._session.flush()
        await self._audit.record(
            AuditAction.IMPORT_ROLLBACK,
            entity_type="import_job",
            entity_id=job.id,
            changes={"rolled_back_at": {"old": None, "new": job.rolled_back_at.isoformat()}},
        )
        return True
