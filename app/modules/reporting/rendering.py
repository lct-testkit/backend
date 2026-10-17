"""Рендеринг отчётов: xlsx / pdf / png (new_spec §4.13, §2.3, §6).

* **xlsx** — `openpyxl` в `write_only`-режиме (раздел 4.13: «стриминг,
  константная память») — тот же приём, что `imports.service` уже применяет к
  сгенерированным отчётам об ошибках импорта. Исключение по форме, не по режиму — файл для LMS
  (`lms_users_upload`): он повторяет шаблон заказчика, см. `render_lms_users_xlsx`.
* **pdf** — раздел 2.3 называет WeasyPrint, но `signing.rendering` (спринт 6)
  уже осознанно заменил его на `xhtml2pdf` (WeasyPrint тянет системные
  Pango/Cairo/GDK-Pixbuf, что противоречит принципу `Dockerfile`) и вложил
  шрифт DejaVu Sans для кириллицы. Здесь этот же пайплайн переиспользуется
  напрямую (`render_template_html`/`html_to_pdf`) — один Jinja2-шаблон
  таблицы на все виды отчёта: колонки/строки в отчёте динамические, а не
  поля с фиксированными именами, поэтому шаблон параметризован структурой
  `ReportDataset`, а не конкретным видом отчёта.
* **png** — раздел 2.3/4.13: «серверный matplotlib (не скриншот фронта — на
  сервере нет браузера)», backend `Agg` — без дисплея. matplotlib поставляет
  свой DejaVu Sans (та же гарнитура, что уже вложена для PDF) и уже
  поддерживает кириллицу без дополнительных шрифтов. Графики есть только для
  видов, где они осмысленны (воронка, динамика по месяцам) — остальные виды
  отчёта отдают только табличные форматы (`report_templates.output_formats`
  ограничивает выбор на уровне API, раздел 4.13 не требует графика для
  каждого отчёта).
"""

from __future__ import annotations

import datetime as dt
import io
import json
import uuid
import zoneinfo
from decimal import Decimal
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402 — backend должен быть выбран до импорта pyplot
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.modules.catalog import learner
from app.modules.imports.parsing import sanitize_formula
from app.modules.reporting.builders import (
    DEFAULT_TIMEZONE_NAME,
    LMS_USERS_UPLOAD,
    ReportDataset,
)
from app.modules.reporting.models import ReportFormat
from app.modules.reporting.rendering_xls import ReportTooLargeError, render_xls, sheet_title
from app.modules.reporting.spec import derived_chart_kinds
from app.modules.signing.rendering import html_to_pdf, render_template_html

__all__ = [
    "CHART_KINDS",
    "CONTENT_TYPES",
    "FILE_EXTENSIONS",
    "RenderNotSupportedError",
    "ReportTooLargeError",
    "render_json",
    "render_lms_users_xlsx",
    "render_pdf",
    "render_png",
    "render_report",
    "render_xls",
    "render_xlsx",
]

CONTENT_TYPES: dict[str, str] = {
    ReportFormat.XLSX.value: ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    # B1-4: `application/vnd.ms-excel` — тип содержимого классического BIFF8/OLE2, а не xlsx.
    ReportFormat.XLS.value: "application/vnd.ms-excel",
    ReportFormat.PDF.value: "application/pdf",
    ReportFormat.PNG.value: "image/png",
    ReportFormat.JSON.value: "application/json",
}

#: Расширение файла по формату — B1-4/B1-5: раньше имя файла строилось как `f"{title}.{job.format}"`
#: (`reporting.service.generate`), что совпадало с форматом случайно (xlsx/pdf/png — расширение
#: и код формата совпадают); xls и json тоже совпадают буквально, но словарь делает это явным и
#: единообразным местом правки, если когда-нибудь появится формат, где это не так.
FILE_EXTENSIONS: dict[str, str] = {fmt.value: fmt.value for fmt in ReportFormat}

#: Виды отчёта, для которых определён график. Ключ — `report_templates.code`. B1-2/B1-8:
#: производный от `ReportSpec.chart` (`reporting.spec`), а не самостоятельный литерал — иначе
#: новый график в `ReportSpec` и регистрация в `rendering.render_png` могли бы расползтись.
CHART_KINDS = derived_chart_kinds()


class RenderNotSupportedError(Exception):
    """PNG запрошен для вида отчёта без определённого графика.

    Не пользовательская ошибка в обычном смысле: `reporting.service`
    проверяет `format in template.output_formats` до вызова рендеринга,
    поэтому попадание сюда означает рассинхрон между посевом шаблонов и
    `CHART_KINDS`, а не некорректный ввод пользователя.
    """


_TABLE_TEMPLATE = """
<h2>{{ title }}</h2>
{% if note %}<p><em>{{ note }}</em></p>{% endif %}
<table style="width:100%; border-collapse:collapse;">
  <thead>
    <tr>
      {% for col in columns %}
      <th style="border:1px solid #999; padding:4px 8px; background:#eef1fa; text-align:left;">
        {{ col }}
      </th>
      {% endfor %}
    </tr>
  </thead>
  <tbody>
    {% for row in rows %}
    <tr>
      {% for cell in row %}
      <td style="border:1px solid #ccc; padding:4px 8px;">{{ "" if cell is none else cell }}</td>
      {% endfor %}
    </tr>
    {% endfor %}
  </tbody>
</table>
<p style="color:#666; font-size:9px;">Сформировано: {{ generated_at }}</p>
"""


def _xlsx_cell(value: Any) -> Any:
    """Раздел 4.12 «запрет формул» / CSV-инъекция: значения отчёта в
    конечном счёте происходят из свободных текстовых полей (название
    организации, ФИО КАМа, название сделки), которые пользователь мог
    ввести с ведущим `=`/`+`/`-`/`@`. `imports.parsing.sanitize_formula` уже
    решает это для сгенерированных xlsx импорта — тот же риск, тот же приём."""
    if value is None:
        return ""
    if isinstance(value, str):
        return sanitize_formula(value)
    if isinstance(value, int | float | bool):
        return value
    return str(value)


def render_xlsx(dataset: ReportDataset) -> bytes:
    workbook = Workbook(write_only=True)
    # B1-4: `sheet_title` — общая с xls функция (`rendering_xls`), чистит запрещённые Excel
    # знаки (`[]:*?/\`) из названия отчёта, не только обрезает длину до 31 знака, как было.
    sheet = workbook.create_sheet(title=sheet_title(dataset.title))
    sheet.append(dataset.columns)
    for row in dataset.rows:
        sheet.append([_xlsx_cell(value) for value in row])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# Ширина колонок листа «Лист1» шаблона LMS — как в оригинале заказчика («Загрузка пользователей»),
# по порядку колонок `learner.LMS_USER_COLUMNS`. «Пол» (L) в оригинале без своей ширины — остаётся
# стандартной.
_LMS_COLUMN_WIDTHS: tuple[float | None, ...] = (
    23.86,  # A  Фамилия
    24.86,  # B  Имя
    24.0,  # C  Отчество
    22.14,  # D  Номер телефона
    20.71,  # E  Email
    14.71,  # F  СНИЛС
    14.57,  # G  Серия паспорта
    15.86,  # H  Номер паспорта
    19.57,  # I  Кем выдан паспорт
    20.43,  # J  Дата выдачи паспорта
    18.43,  # K  Код подразделения
    None,  # L  Пол
    15.29,  # M  Дата рождения
    18.14,  # N  Регион регистрации
    29.14,  # O  Населенный пункт регистрации
    17.71,  # P  Улица регистрации
    16.29,  # Q  Дом регистрации
    20.29,  # R  Квартира регистрации
    19.0,  # S  Индекс регистрации
    23.0,  # T  Имя (дательный падеж)
    27.0,  # U  Фамилия (дательный падеж)
    26.14,  # V  Отчество (дательный падеж)
    37.14,  # W  Образование
    21.43,  # X  Профессия по диплому
    29.57,  # Y  Учебное заведение по диплому
    29.0,  # Z  Фамилия, указанная в дипломе
    15.29,  # AA Номер диплома
    14.86,  # AB Серия диплома
    31.0,  # AC Регистрационный номер диплома
    22.0,  # AD Дата выдачи диплома
)
_LMS_LOOKUP_WIDTH = 58.14
_LMS_DATE_FORMAT = "DD.MM.YYYY"
# В оригинале первые шесть заголовков выровнены по центру, остальные — по умолчанию.
_LMS_CENTERED_HEADERS = 6


def _lms_cell(sheet: Any, value: Any) -> Any:
    """Ячейка данных шаблона LMS. Число (телефон `79990234365`) — числом, дата — настоящей датой
    Excel в формате `ДД.ММ.ГГГГ`, всё остальное — текстом (СНИЛС, серия и номер паспорта, индекс:
    иначе Excel съел бы ведущие нули и дефисы). Пустое остаётся пустой ячейкой.

    Текст проходит через `sanitize_formula`: значения приходят от людей (адрес, ФИО, название
    учебного заведения), и ячейка, начинающаяся с `=`, стала бы формулой — файл заказчик открывает в
    Excel. Формул в шаблоне нет."""
    if value is None or value == "":
        return None
    if isinstance(value, dt.date):
        cell = WriteOnlyCell(sheet, value=value)
        cell.number_format = _LMS_DATE_FORMAT
        return cell
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return value
    return sanitize_formula(str(value))


def render_lms_users_xlsx(dataset: ReportDataset) -> bytes:
    """Файл для загрузки учащихся в LMS: повторяет шаблон заказчика `Загрузка пользователей.xlsx`.

    * `Лист1` — 30 заголовков буквально как в оригинале (в том числе `Отчествопри наличии)`: LMS
      сопоставляет колонки по этому тексту) и строки учащихся; ширина колонок как в оригинале.
    * `Лист2` — справочники: `М`/`Ж` в A1:A2 и семь уровней образования в B1:B7.
    * На «Пол» (L) и «Образование» (W) повешены выпадающие списки на этот справочник, на 1000 строк.

    Книга пишется в `write_only`-режиме, как остальные отчёты: константная память, а выпадающие
    списки, ширина колонок и второй лист в нём поддерживаются (ширина задаётся до записи первой
    строки листа, списки — до сохранения книги).
    """
    workbook = Workbook(write_only=True)
    users = workbook.create_sheet(title=learner.LMS_USERS_SHEET)
    lookup = workbook.create_sheet(title=learner.LMS_LOOKUP_SHEET)

    for index, width in enumerate(_LMS_COLUMN_WIDTHS, start=1):
        if width is not None:
            users.column_dimensions[get_column_letter(index)].width = width
    lookup.column_dimensions["B"].width = _LMS_LOOKUP_WIDTH

    header_row = []
    for index, title in enumerate(dataset.columns):
        header = WriteOnlyCell(users, value=title)
        header.font = Font(name="Calibri", size=11, bold=True)
        if index < _LMS_CENTERED_HEADERS:
            header.alignment = Alignment(horizontal="center")
        header_row.append(header)
    users.append(header_row)
    for row in dataset.rows:
        users.append([_lms_cell(users, value) for value in row])

    sex_labels = list(learner.SEX_LABELS.values())
    education_labels = [label for _code, label in learner.EDUCATION_LEVELS]
    for position in range(max(len(sex_labels), len(education_labels))):
        lookup.append(
            [
                sex_labels[position] if position < len(sex_labels) else None,
                education_labels[position] if position < len(education_labels) else None,
            ]
        )

    targets = [target for target, _header in learner.LMS_USER_COLUMNS]
    for target, source in (
        ("sex", f"{learner.LMS_LOOKUP_SHEET}!$A$1:$A${len(sex_labels)}"),
        ("education", f"{learner.LMS_LOOKUP_SHEET}!$B$1:$B${len(education_labels)}"),
    ):
        validation = DataValidation(
            type="list", formula1=source, allow_blank=True, showErrorMessage=True
        )
        column = get_column_letter(targets.index(target) + 1)
        validation.add(f"{column}2:{column}{learner.LMS_VALIDATION_LAST_ROW}")
        # У write-only листа нет `add_data_validation`, но список проверок в нём есть и
        # сохраняется вместе с книгой.
        users.data_validations.append(validation)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def render_pdf(dataset: ReportDataset) -> bytes:
    context = {
        "title": dataset.title,
        "note": dataset.note,
        "columns": dataset.columns,
        "rows": dataset.rows,
        "generated_at": dataset.generated_at.strftime("%d.%m.%Y %H:%M UTC"),
    }
    return html_to_pdf(render_template_html(_TABLE_TEMPLATE, context))


#: rtk_requiriments.md §6 п.4: «возможность формирования результирующего json-файла».
#: Версионировано — пересматривается при несовместимых изменениях структуры файла (B1-5).
_JSON_SCHEMA = "rtk-crm-report/1"


def _json_cell(value: Any, tz: dt.tzinfo) -> Any:
    """Значение ячейки для json-файла (B1-5): без формул/усечения строк (`/data` и файл отдают
    данные «как есть», в отличие от xlsx/xls, где `sanitize_formula` защищает от CSV-инъекции —
    тот риск специфичен открытию в Excel, не парсингу JSON). `Decimal` — числом, aware-`datetime`
    переводится в пояс пользователя и остаётся со смещением (`2026-01-05T10:30:00+03:00`), дата
    без времени — `YYYY-MM-DD`."""
    if value is None or isinstance(value, bool | int | str):
        return value
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, dt.datetime):
        local = value.astimezone(tz) if value.tzinfo is not None else value
        return local.isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float):
        return value
    return str(value)


def render_json(
    dataset: ReportDataset,
    *,
    template_code: str,
    job_id: uuid.UUID | None,
    params: dict[str, Any] | None,
    tz_name: str | None,
) -> bytes:
    """Файл для `POST /api/reports` с `format=json` (B1-5). Тот же `ReportDataset`, что рендерится
    в xlsx/pdf/png, версионированным конвертом: `columns` всегда идут вместе со `rows` (B1-2:
    `dataset.keys`/`dataset.kinds`, пусто у видов без каталога колонок — тогда `key` у каждой
    колонки `None`), а не массивами по позиции без заголовков, — так перестановка колонок не
    ломает интегратора, который уже читает файл по ключам. Кодировка UTF-8 без BOM (`ensure_
    ascii=False`), `allow_nan=False` — `float('nan')`/`inf` в данных поднимет понятную `ValueError`
    при сериализации, а не отдаст невалидный JSON наружу."""
    try:
        tz = zoneinfo.ZoneInfo(tz_name) if tz_name else zoneinfo.ZoneInfo(DEFAULT_TIMEZONE_NAME)
    except (zoneinfo.ZoneInfoNotFoundError, ValueError):
        tz = zoneinfo.ZoneInfo(DEFAULT_TIMEZONE_NAME)

    keys = dataset.keys if dataset.keys is not None else [None] * len(dataset.columns)
    kinds = dataset.kinds if dataset.kinds is not None else ["text"] * len(dataset.columns)
    columns = [
        {"key": key, "title": title, "kind": kind}
        for key, title, kind in zip(keys, dataset.columns, kinds, strict=True)
    ]
    # Служебные ключи (`_seed` и т.п., B1-2) не предназначены для интегратора.
    visible_params = {k: v for k, v in (params or {}).items() if not str(k).startswith("_")}

    payload = {
        "schema": _JSON_SCHEMA,
        "report": {
            "code": template_code,
            "title": dataset.title,
            "job_id": str(job_id) if job_id is not None else None,
            "generated_at": dataset.generated_at.astimezone(tz).isoformat(),
            "timezone": tz_name or DEFAULT_TIMEZONE_NAME,
        },
        "params": visible_params,
        "columns": columns,
        "rows": [[_json_cell(value, tz) for value in row] for row in dataset.rows],
        "totals": {
            "row_count": len(dataset.rows),
            "total_rows": (
                dataset.total_rows if dataset.total_rows is not None else len(dataset.rows)
            ),
            "truncated": dataset.truncated,
        },
        "note": dataset.note,
    }
    return json.dumps(payload, ensure_ascii=False, allow_nan=False).encode("utf-8")


#: Палитра и marks-спеки — из внутренней методики визуализации (skill
#: `dataviz`, `references/palette.md`/`marks-and-anatomy.md`), уже
#: провалидированной её собственным скриптом (`validate_palette.js`), а не
#: подобранной на глаз. Раньше здесь был один плоский `#4C6EF5` на воронку
#: из 14+ статусов и цикл matplotlib по умолчанию на линиях — оба варианта
#: читаемы, только когда категорий мало; на реальной 14-шаговой B2B-воронке
#: (раздел 1) подписи статусов друг на друга налезали (вертикальные бары +
#: поворот 25°), это и была жалоба, из-за которой переписан весь модуль.
_SURFACE = "#fcfcfb"
_INK_PRIMARY = "#0b0b0b"
_INK_SECONDARY = "#52514e"
_INK_MUTED = "#898781"
_GRIDLINE = "#e1e0d9"
_BASELINE = "#c3c2b7"

#: Ordinal-рамп (позиция в последовательности — «funnel stage» это
#: буквальный пример ordinal-job в color-formula.md): один тон, монотонная
#: светлота, светлый конец не темнее шага 250 (порог контраста ordinal-
#: рампы на светлой поверхности, см. palette.md).
_ORDINAL_BLUE_STEPS = [
    "#86b6ef",
    "#6da7ec",
    "#5598e7",
    "#3987e5",
    "#2a78d6",
    "#256abf",
    "#1c5cab",
    "#184f95",
    "#104281",
    "#0d366b",
]

#: Категориальные слоты 1/2/3 в фиксированном порядке (palette.md) —
#: никогда не переставляются и не генерируются заново, это и есть механизм
#: CVD-безопасности (color-formula.md, «Fixed hue anchors»).
_CATEGORICAL = {"blue": "#2a78d6", "orange": "#eb6834", "aqua": "#1baf7a"}


def _ordinal_ramp(n: int) -> list[str]:
    """N цветов, равномерно растянутых по `_ORDINAL_BLUE_STEPS`."""
    if n <= 1:
        return [_ORDINAL_BLUE_STEPS[len(_ORDINAL_BLUE_STEPS) // 2]]
    last = len(_ORDINAL_BLUE_STEPS) - 1
    return [_ORDINAL_BLUE_STEPS[round(i * last / (n - 1))] for i in range(n)]


def _format_value(value: float) -> str:
    """Табличные числа с разрядным пробелом (marks-and-anatomy.md: «round
    to clean numbers... thousands-comma'd» — пробел, не запятая: русская
    типографская норма для разрядов, раздел 2 «денежные суммы» уже
    использует ту же логику для сумм)."""
    return f"{value:,.0f}".replace(",", " ")


def _style_axes(ax: Any) -> None:
    """Общий «тихий» chrome (marks-and-anatomy.md): убираем рамку графика,
    оставляем только нижнюю ось базовой линией, подписи — приглушённым
    тоном, а не чёрным по умолчанию."""
    for side in ("top", "right", "left"):
        ax.spines[side].set_visible(False)
    ax.spines["bottom"].set_color(_BASELINE)
    ax.tick_params(colors=_INK_MUTED, labelsize=9)


def _fig_to_png(fig: Any) -> bytes:
    buffer = io.BytesIO()
    fig.patch.set_facecolor(_SURFACE)
    fig.savefig(buffer, format="png", dpi=130, bbox_inches="tight", facecolor=_SURFACE)
    plt.close(fig)
    return buffer.getvalue()


def _render_funnel_chart(dataset: ReportDataset) -> bytes:
    # Горизонтальные бары, не вертикальные: раздел 1 описывает воронку до
    # 14 шагов с длинными названиями статусов («Юридическое согласование
    # договора», «Закрытие периода и пролонгация») — на вертикальных барах
    # это решалось поворотом подписи, который и порождал налезающий текст.
    # На горизонтальных барах подпись читается вдоль своей оси без поворота
    # при любой длине названия и любом числе статусов.
    labels = [str(row[0]) for row in dataset.rows]
    currently_in = [float(row[1] or 0) for row in dataset.rows]
    colors = _ordinal_ramp(len(labels))

    # barh рисует снизу вверх; разворачиваем все три параллельных списка ОДИН
    # раз здесь, чтобы первый шаг воронки (раздел 1) оказался вверху графика,
    # как читатель ожидает — а дальше индексация везде простая, без второго
    # слоя реверса. (Раньше реверс применялся только к позициям баров, но не
    # синхронно к подписям значений — бар получал верное место, а число рядом
    # с ним оказывалось от совсем другого статуса; поймано на рендере, не на
    # чтении кода — ровно то, для чего в методике есть шаг «отрендери и
    # посмотри».)
    labels = labels[::-1]
    currently_in = currently_in[::-1]
    colors = colors[::-1]

    # Высота фигуры растёт вместе с числом статусов — раньше она была
    # фиксированной (4.5 дюйма) независимо от того, 6 строк в воронке или 17.
    height = max(3.2, 0.5 * len(labels) + 1.2)
    fig, ax = plt.subplots(figsize=(8.5, height))

    y_pos = list(range(len(labels)))
    ax.barh(
        y_pos,
        currently_in,
        height=0.62,
        color=colors,
        edgecolor=_SURFACE,
        linewidth=2,  # 2px surface-gap между соседними барами
    )
    ax.set_yticks(y_pos)
    ax.set_yticklabels(labels, color=_INK_SECONDARY)

    max_value = max(currently_in, default=0)
    ax.set_xlim(0, max_value * 1.18 if max_value else 1)
    # Bars → value at the tip (marks-and-anatomy.md) — по одному числу на
    # бар, это не «число на каждой точке» (то правило — про линии/точки).
    for y, value in zip(y_pos, currently_in, strict=True):
        ax.text(
            value + max_value * 0.02,
            y,
            _format_value(value),
            va="center",
            ha="left",
            fontsize=9,
            color=_INK_PRIMARY,
        )

    ax.set_title(dataset.title, color=_INK_PRIMARY, fontsize=13, loc="left", pad=12)
    ax.set_xlabel("Сделок сейчас в статусе", color=_INK_SECONDARY, fontsize=9)
    ax.grid(axis="x", color=_GRIDLINE, linewidth=1, linestyle="-", zorder=0)
    ax.set_axisbelow(True)
    _style_axes(ax)
    ax.spines["bottom"].set_visible(False)
    fig.tight_layout()
    return _fig_to_png(fig)


def _render_monthly_chart(dataset: ReportDataset) -> bytes:
    months = [str(row[0]) for row in dataset.rows]
    series = [
        ("Создано", [float(row[1] or 0) for row in dataset.rows], _CATEGORICAL["blue"]),
        ("Выиграно", [float(row[2] or 0) for row in dataset.rows], _CATEGORICAL["orange"]),
        ("Проиграно", [float(row[3] or 0) for row in dataset.rows], _CATEGORICAL["aqua"]),
    ]
    fig, ax = plt.subplots(figsize=(8.5, 4.6))
    x = range(len(months))

    for name, values, color in series:
        ax.plot(
            x,
            values,
            color=color,
            linewidth=2,
            marker="o",
            markersize=8,
            markerfacecolor=color,
            markeredgecolor=_SURFACE,
            markeredgewidth=2,
            label=name,
            zorder=3,
        )
        # Direct end-label (marks-and-anatomy.md: «Lines → value at the
        # end») — до 4 серий подписываются и напрямую, легенда не
        # единственный способ понять, какая линия какая. Текст — тоном
        # чернил, не цветом серии («Text never wears the data color»):
        # идентичность несёт цветной маркер рядом, не сама подпись.
        if values:
            ax.annotate(
                f"{name}: {_format_value(values[-1])}",
                (x[-1], values[-1]),
                xytext=(8, 0),
                textcoords="offset points",
                va="center",
                fontsize=9,
                color=_INK_SECONDARY,
            )

    ax.set_title(dataset.title, color=_INK_PRIMARY, fontsize=13, loc="left", pad=12)
    ax.set_xticks(list(x))
    # Короткие подписи месяцев (YYYY-MM) — поворот только когда их много
    # настолько, что горизонтально они начинают соприкасаться; раньше был
    # фиксированный поворот 40° даже на 3-4 месяца, где он не нужен.
    rotation = 30 if len(months) > 8 else 0
    ax.set_xticklabels(months, rotation=rotation, ha="right" if rotation else "center")
    ax.grid(axis="y", color=_GRIDLINE, linewidth=1, linestyle="-", zorder=0)
    ax.set_axisbelow(True)
    _style_axes(ax)
    legend = ax.legend(
        loc="upper left",
        frameon=False,
        fontsize=9,
        labelcolor=_INK_SECONDARY,
    )
    for handle in legend.legend_handles:
        handle.set_markeredgecolor(_SURFACE)
    fig.tight_layout()
    return _fig_to_png(fig)


def render_png(dataset: ReportDataset, *, template_code: str) -> bytes:
    if template_code == "deal_funnel":
        return _render_funnel_chart(dataset)
    if template_code == "monthly_dynamics":
        return _render_monthly_chart(dataset)
    raise RenderNotSupportedError(f"Для вида отчёта {template_code!r} график не определён")


def render_report(
    dataset: ReportDataset,
    *,
    format: str,
    template_code: str,
    job_id: uuid.UUID | None = None,
    params: dict[str, Any] | None = None,
    tz_name: str | None = None,
) -> bytes:
    """`job_id`/`params`/`tz_name` — только для `format=json` (B1-5): конверт файла несёт, с
    какими параметрами и для какого задания он построен. Остальные форматы их не используют —
    они приходят через один и тот же вызов из `reporting.service._render_serialized`, чтобы не
    заводить отдельную функцию рендера ради трёх новых аргументов одного формата из пяти."""
    if format == ReportFormat.XLSX.value:
        if template_code == LMS_USERS_UPLOAD:
            return render_lms_users_xlsx(dataset)
        return render_xlsx(dataset)
    if format == ReportFormat.XLS.value:
        return render_xls(dataset)
    if format == ReportFormat.PDF.value:
        return render_pdf(dataset)
    if format == ReportFormat.PNG.value:
        return render_png(dataset, template_code=template_code)
    if format == ReportFormat.JSON.value:
        return render_json(
            dataset, template_code=template_code, job_id=job_id, params=params, tz_name=tz_name
        )
    raise RenderNotSupportedError(f"Неизвестный формат отчёта {format!r}")
