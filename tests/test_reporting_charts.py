"""Тесты диаграмм отчётов (B1-8, спринт B1): `reporting.charts` — построение `Figure` для пяти
видов графика (воронка, динамика, рейтинг востребованности, прогресс обучения, реестр лицензий)
и встраивание того же PNG в PDF (первая страница) и xlsx (лист «График»). Без БД — чистые
функции над готовым `ReportDataset`, тот же приём, что `tests/test_reporting.py::TestPngRendering`
уже применяет к воронке/динамике.

Объективность проверяется по объекту `Figure`, а не по самой картинке: ширина/высота баров,
`ax.get_xlim()[0] == 0`/`ax.get_ylim()[0] == 0`, подпись значения у каждого бара, текст «Нет
данных за выбранный период» на пустом наборе."""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app.modules.reporting.builders import ReportDataset
from app.modules.reporting.charts import (
    CHART_BUILDERS,
    RenderNotSupportedError,
    build_demand_figure,
    build_funnel_figure,
    build_license_status_figure,
    build_monthly_figure,
    build_progress_figure,
    figure_to_png,
    progress_buckets,
)
from app.modules.reporting.charts import (
    render_png as charts_render_png,
)
from app.modules.reporting.rendering import CHART_KINDS, render_pdf, render_png, render_xlsx
from app.modules.reporting.spec import REPORT_SPECS


def _funnel_dataset(rows: list[list] | None = None) -> ReportDataset:
    return ReportDataset(
        title="Воронка по статусам",
        columns=[
            "Статус",
            "Сейчас в статусе",
            "Всего прошло",
            "Конверсия шага, %",
            "Среднее время, дн.",
            "Дошли до шага",
            "От старта, %",
        ],
        rows=rows
        if rows is not None
        else [
            ["Новый", 10, 100, 100.0, 1.0, 100, 100.0],
            ["В работе", 5, 83, 83.0, 2.0, 83, 83.0],
            ["Выиграно", 3, 40, 48.2, 5.0, 40, 40.0],
            ["Отказ", 12, 12, None, None, None, 12.0],
            ["Заморожена", 4, 5, None, None, None, 5.0],
        ],
    )


def _demand_dataset(n: int = 5) -> ReportDataset:
    rows = [
        [i + 1, f"Продукт {i + 1}", "Направление A", 40 - i, 10, 300 - i * 5, 3, 50.0 - i]
        for i in range(n)
    ]
    return ReportDataset(
        title="Рейтинг востребованности ИТ-продуктов и направлений",
        columns=[
            "Место",
            "Продукт",
            "Направление",
            "Заявок",
            "Выиграно",
            "Обучающихся",
            "Потоков",
            "Индекс, %",
        ],
        rows=rows,
    )


def _progress_dataset(values: list[int] | None = None) -> ReportDataset:
    values = values if values is not None else [0, 25, 26, 100, 80, 45]
    return ReportDataset(
        title="Прогресс обучения",
        columns=["Студент", "Продукт", "Прогресс, %"],
        rows=[["Учащийся", "P1", v] for v in values],
        extra={"progress_values": values},
    )


def _license_status_dataset(
    by_status: dict[str, int] | None = None, by_year: dict[str, int] | None = None
) -> ReportDataset:
    return ReportDataset(
        title="Реестр лицензий",
        columns=["Вуз", "Вендор"],
        rows=[["МГТУ", "Яндекс"]],
        extra={
            "by_status": by_status if by_status is not None else {"В работе": 3, "Передано": 5},
            "by_year": by_year if by_year is not None else {"2026": 4, "2027": 5},
        },
    )


class TestChartKindsRegistry:
    def test_chart_builders_cover_every_registered_chart_kind(self) -> None:
        """Каждый `ReportSpec.chart` должен иметь функцию построения — иначе `render_png`
        рассинхронизирован с сидом/`REPORT_SPECS` (та же защита, что `test_reporting.
        py::test_png_only_seeded_for_kinds_with_a_registered_chart` проверяет с другой стороны)."""
        chart_kinds = {spec.chart for spec in REPORT_SPECS.values() if spec.chart}
        assert chart_kinds <= set(CHART_BUILDERS)

    def test_five_report_kinds_have_a_chart(self) -> None:
        assert {
            "deal_funnel",
            "monthly_dynamics",
            "product_demand",
            "learning_progress",
            "license_register",
        } == CHART_KINDS

    def test_unknown_chart_kind_raises(self) -> None:
        with pytest.raises(RenderNotSupportedError):
            charts_render_png(_funnel_dataset(), chart_kind="unknown")


class TestFunnelFigure:
    def test_bar_widths_match_passed_and_currently_in(self) -> None:
        fig = build_funnel_figure(_funnel_dataset())
        ax = fig.axes[0]
        # Два вызова `barh` — «Всего прошло» (полупрозрачный) и «Сейчас в статусе» (сверху);
        # bars идут в порядке добавления, первые 5 патчей — светлый слой, следующие 5 — тёмный.
        widths = [p.get_width() for p in ax.patches]
        assert len(widths) == 10
        passed = sorted(widths[:5], reverse=True)
        assert passed == [100, 83, 40, 12, 5]

    def test_x_axis_starts_at_zero(self) -> None:
        fig = build_funnel_figure(_funnel_dataset())
        assert fig.axes[0].get_xlim()[0] == 0

    def test_every_bar_has_a_value_label(self) -> None:
        fig = build_funnel_figure(_funnel_dataset())
        ax = fig.axes[0]
        assert len(ax.texts) == 5

    def test_main_step_label_includes_conversion(self) -> None:
        fig = build_funnel_figure(_funnel_dataset())
        texts = [t.get_text() for t in fig.axes[0].texts]
        assert any("83" in t and "от предыдущего" in t for t in texts)

    def test_lost_and_parked_labels_have_no_conversion_phrase(self) -> None:
        fig = build_funnel_figure(_funnel_dataset())
        texts = [t.get_text() for t in fig.axes[0].texts]
        # «Отказ»/«Заморожена» — значения 12 и 5, без «от предыдущего» рядом с этими числами.
        assert any(t == "12" for t in texts)
        assert any(t == "5" for t in texts)

    def test_lost_parked_group_is_separated_by_a_line(self) -> None:
        fig = build_funnel_figure(_funnel_dataset())
        assert len(fig.axes[0].lines) == 1

    def test_no_data_message_on_empty_rows(self) -> None:
        fig = build_funnel_figure(_funnel_dataset(rows=[]))
        texts = [t.get_text() for t in fig.axes[0].texts]
        assert any("Нет данных" in t for t in texts)

    def test_short_legacy_rows_do_not_crash(self) -> None:
        """Контракт билдера (B1-7) — 7 колонок, но старые/минимальные тестовые строки короче
        (`tests/test_reporting.py::TestPngRendering`); график не должен падать `IndexError`."""
        fig = build_funnel_figure(_funnel_dataset(rows=[["Новый", 10], ["В работе", 5]]))
        assert len(fig.axes[0].patches) == 4


class TestMonthlyFigure:
    def test_y_axis_starts_at_zero(self) -> None:
        dataset = ReportDataset(
            title="Динамика",
            columns=["Месяц", "Создано", "Выиграно", "Проиграно"],
            rows=[["2026-01", 5, 2, 1], ["2026-02", 0, 0, 0], ["2026-03", 8, 4, 2]],
        )
        fig = build_monthly_figure(dataset)
        assert fig.axes[0].get_ylim()[0] == 0

    def test_no_data_message_on_empty_rows(self) -> None:
        dataset = ReportDataset(title="Динамика", columns=["Месяц"], rows=[])
        fig = build_monthly_figure(dataset)
        texts = [t.get_text() for t in fig.axes[0].texts]
        assert any("Нет данных" in t for t in texts)


class TestDemandFigure:
    def test_three_panels_with_independent_zero_based_scales(self) -> None:
        fig = build_demand_figure(_demand_dataset())
        assert len(fig.axes) == 3
        for ax in fig.axes:
            assert ax.get_xlim()[0] == 0

    def test_bar_widths_match_applications_learners_streams(self) -> None:
        fig = build_demand_figure(_demand_dataset(n=3))
        applications_ax, learners_ax, streams_ax = fig.axes
        assert sorted(p.get_width() for p in applications_ax.patches) == [38, 39, 40]
        assert sorted(p.get_width() for p in learners_ax.patches) == [290, 295, 300]
        assert [p.get_width() for p in streams_ax.patches] == [3, 3, 3]

    def test_truncates_to_first_fifteen_and_notes_it(self) -> None:
        fig = build_demand_figure(_demand_dataset(n=20))
        assert len(fig.axes[0].patches) == 15
        captions = [t.get_text() for t in fig.texts]
        assert any("первые 15 из 20" in c for c in captions)

    def test_no_data_message_on_empty_rows(self) -> None:
        fig = build_demand_figure(_demand_dataset(n=0))
        texts = [t.get_text() for t in fig.axes[0].texts]
        assert any("Нет данных" in t for t in texts)


class TestProgressBuckets:
    def test_pure_function_sorts_into_six_buckets(self) -> None:
        buckets = progress_buckets([0, 1, 25, 26, 50, 51, 75, 76, 99, 100])
        counts = dict(buckets)
        assert counts["0%"] == 1
        assert counts["1–25%"] == 2
        assert counts["26–50%"] == 2
        assert counts["51–75%"] == 2
        assert counts["76–99%"] == 2
        assert counts["100%"] == 1

    def test_empty_input_gives_zero_counts(self) -> None:
        buckets = progress_buckets([])
        assert all(count == 0 for _label, count in buckets)


class TestProgressFigure:
    def test_bar_heights_match_bucket_counts(self) -> None:
        # 0→"0%", 25→"1–25%", 26 и 45→"26–50%", 80→"76–99%", 100→"100%": [1,1,2,0,1,1].
        fig = build_progress_figure(_progress_dataset([0, 25, 26, 100, 80, 45]))
        heights = sorted(p.get_height() for p in fig.axes[0].patches)
        assert heights == [0, 1, 1, 1, 1, 2]

    def test_caption_reports_count_and_average(self) -> None:
        fig = build_progress_figure(_progress_dataset([0, 100]))
        captions = [t.get_text() for t in fig.texts]
        assert any("n = 2 записей" in c and "50%" in c for c in captions)

    def test_no_data_message_when_extra_missing(self) -> None:
        dataset = ReportDataset(title="Прогресс обучения", columns=["A"], rows=[])
        fig = build_progress_figure(dataset)
        texts = [t.get_text() for t in fig.axes[0].texts]
        assert any("Нет данных" in t for t in texts)


class TestLicenseStatusFigure:
    def test_two_panels(self) -> None:
        fig = build_license_status_figure(_license_status_dataset())
        assert len(fig.axes) == 2

    def test_bar_heights_match_counts(self) -> None:
        fig = build_license_status_figure(
            _license_status_dataset(by_status={"В работе": 3, "Передано": 5})
        )
        heights = sorted(p.get_height() for p in fig.axes[0].patches)
        assert heights == [3, 5]

    def test_no_data_message_when_extra_empty(self) -> None:
        dataset = _license_status_dataset(by_status={}, by_year={})
        fig = build_license_status_figure(dataset)
        texts = [t.get_text() for t in fig.axes[0].texts]
        assert any("Нет данных" in t for t in texts)


class TestPngMagicBytes:
    @pytest.mark.parametrize(
        ("template_code", "dataset_factory"),
        [
            ("product_demand", _demand_dataset),
            ("learning_progress", _progress_dataset),
            ("license_register", _license_status_dataset),
        ],
    )
    def test_png_starts_with_magic_bytes(self, template_code, dataset_factory) -> None:
        content = render_png(dataset_factory(), template_code=template_code)
        assert content.startswith(b"\x89PNG\r\n\x1a\n")

    def test_figure_to_png_without_pyplot_state(self) -> None:
        """`figure_to_png` строит PNG из `Figure`, созданной напрямую (без `pyplot`) — повторный
        вызов не должен ничего накапливать в глобальном реестре (которого здесь просто нет)."""
        fig = build_monthly_figure(
            ReportDataset(title="Д", columns=["M"], rows=[["2026-01", 1, 0, 0]])
        )
        png = figure_to_png(fig)
        assert png.startswith(b"\x89PNG\r\n\x1a\n")


class TestPdfChartEmbedding:
    def test_pdf_embeds_one_image_on_first_page_for_chart_template(self) -> None:
        content = render_pdf(_demand_dataset(), template_code="product_demand")
        assert content.startswith(b"%PDF")
        reader = PdfReader(io.BytesIO(content))
        assert len(reader.pages[0].images) == 1

    def test_pdf_without_template_code_has_no_image(self) -> None:
        content = render_pdf(_demand_dataset())
        reader = PdfReader(io.BytesIO(content))
        assert len(reader.pages[0].images) == 0

    def test_pdf_for_template_without_chart_has_no_image(self) -> None:
        dataset = ReportDataset(title="Сводка", columns=["A"], rows=[["x"]])
        content = render_pdf(dataset, template_code="kam_summary")
        reader = PdfReader(io.BytesIO(content))
        assert len(reader.pages[0].images) == 0


class TestXlsxChartSheet:
    def test_xlsx_has_a_chart_sheet_with_one_image_for_chart_template(self) -> None:
        content = render_xlsx(_demand_dataset(), template_code="product_demand")
        workbook = load_workbook(io.BytesIO(content))
        assert "График" in workbook.sheetnames
        assert len(workbook["График"]._images) == 1

    def test_xlsx_without_template_code_has_no_chart_sheet(self) -> None:
        content = render_xlsx(_demand_dataset())
        workbook = load_workbook(io.BytesIO(content))
        assert "График" not in workbook.sheetnames

    def test_xlsx_for_template_without_chart_has_no_chart_sheet(self) -> None:
        dataset = ReportDataset(title="Сводка", columns=["A"], rows=[["x"]])
        content = render_xlsx(dataset, template_code="kam_summary")
        workbook = load_workbook(io.BytesIO(content))
        assert "График" not in workbook.sheetnames
