"""Выгрузка учащихся в шаблон LMS «Загрузка пользователей» (отчёт `lms_users_upload`).

Файл обязан совпасть с шаблоном заказчика: лист «Лист1» с 30 заголовками буквально как в оригинале,
телефон числом, СНИЛС и паспорт текстом, настоящие даты `ДД.ММ.ГГГГ`, лист «Лист2» со справочниками
и выпадающие списки на «Пол» и «Образование». Формул в файле нет. Проверки формы файла — чистые
(`render_lms_users_xlsx`), выборка и весь путь через очередь отчётов — на настоящей Postgres, см.
`tests/conftest.py`.
"""

from __future__ import annotations

import datetime as dt
import io
import uuid
from typing import Any

import openpyxl
import pytest

from app.modules.catalog import learner
from app.modules.reporting.builders import ReportDataset
from app.modules.reporting.rendering import render_lms_users_xlsx, render_report
from tests.conftest import TEST_DATABASE_URL, _make_user, authenticate, run

needs_db = pytest.mark.skipif(
    not TEST_DATABASE_URL, reason="нужен TEST_DATABASE_URL с применёнными миграциями"
)

HEADERS = [header for _target, header in learner.LMS_USER_COLUMNS]
TARGETS = [target for target, _header in learner.LMS_USER_COLUMNS]


def _row(**values: Any) -> list[Any]:
    """Строка датасета в порядке колонок шаблона: то, что отдаёт `build_lms_users_upload`."""
    unknown = set(values) - set(TARGETS)
    assert not unknown, unknown
    return [values.get(target) for target in TARGETS]


def _open(content: bytes) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(content))


def _dataset(*rows: list[Any]) -> ReportDataset:
    return ReportDataset(title="Загрузка пользователей", columns=list(HEADERS), rows=list(rows))


class TestTemplateShape:
    """Чистая часть: форма файла не зависит от БД."""

    def _full_row(self) -> list[Any]:
        return _row(
            last_name="Черепанова",
            first_name="Светлана",
            middle_name="Васильевна",
            phone=79990234365,
            email="cherepanona.s@test.ru",
            snils="112-233-445 95",
            passport_series="4512",
            passport_number="123456",
            passport_issued_by="ОУФМС России по г. Москве",
            passport_issued_at=dt.date(2020, 3, 13),
            passport_dept_code="770-001",
            sex="Ж",
            birth_date=dt.date(1990, 5, 17),
            reg_region="Москва",
            reg_city="Москва",
            reg_street="Тверская",
            reg_house="1",
            reg_apartment="15",
            reg_zip="012345",
            first_name_dative="Светлане",
            last_name_dative="Черепановой",
            middle_name_dative="Васильевне",
            education="Высшее образование – бакалавриат",
            diploma_profession="Программист",
            diploma_institution="Московский государственный университет",
            diploma_surname="Черепанова",
            diploma_number="1234567",
            diploma_series="АБ",
            diploma_reg_number="123456789",
            diploma_issued_at=dt.date(2015, 6, 30),
        )

    def test_two_sheets_with_the_template_names(self) -> None:
        workbook = _open(render_lms_users_xlsx(_dataset()))

        assert workbook.sheetnames == ["Лист1", "Лист2"]

    def test_headers_are_copied_literally_from_the_original(self) -> None:
        sheet = _open(render_lms_users_xlsx(_dataset())).worksheets[0]

        headers = [cell.value for cell in sheet[1]]
        assert headers == HEADERS
        assert len(headers) == 30
        # Скобки потеряны в оригинале, и LMS сопоставляет колонки по этому тексту.
        assert headers[2] == "Отчествопри наличии)"
        assert headers[19] == "Имядательный падеж)"
        assert headers[11] == "Пол"
        assert headers[22] == "Образование"
        assert sheet.max_row == 1

    def test_header_style_matches_the_original(self) -> None:
        sheet = _open(render_lms_users_xlsx(_dataset())).worksheets[0]

        assert all(cell.font.b for cell in sheet[1])
        assert sheet["A1"].alignment.horizontal == "center"
        assert sheet["F1"].alignment.horizontal == "center"
        assert sheet["G1"].alignment.horizontal is None

    def test_column_widths_follow_the_original(self) -> None:
        workbook = _open(render_lms_users_xlsx(_dataset()))

        users = workbook["Лист1"].column_dimensions
        assert users["A"].width == pytest.approx(23.86, abs=0.01)
        assert users["W"].width == pytest.approx(37.14, abs=0.01)
        assert users["AC"].width == pytest.approx(31.0, abs=0.01)
        assert workbook["Лист2"].column_dimensions["B"].width == pytest.approx(58.14, abs=0.01)

    def test_cell_types_follow_the_template(self) -> None:
        sheet = _open(render_lms_users_xlsx(_dataset(self._full_row()))).worksheets[0]

        # Телефон — целое число без плюса, как `79990234365` в оригинале.
        assert sheet["D2"].value == 79990234365
        assert isinstance(sheet["D2"].value, int)
        assert sheet["D2"].data_type == "n"
        # СНИЛС, серия, номер, код подразделения и индекс — текст (нули и дефисы не съедаются).
        for column, expected in (
            ("F", "112-233-445 95"),
            ("G", "4512"),
            ("H", "123456"),
            ("K", "770-001"),
            ("S", "012345"),
        ):
            cell = sheet[f"{column}2"]
            assert cell.value == expected, column
            assert cell.data_type == "s", column
        assert sheet["L2"].value == "Ж"
        assert sheet["W2"].value == "Высшее образование – бакалавриат"
        assert sheet["A2"].value == "Черепанова"
        assert sheet["E2"].value == "cherepanona.s@test.ru"

    def test_dates_are_real_excel_dates_in_the_russian_format(self) -> None:
        sheet = _open(render_lms_users_xlsx(_dataset(self._full_row()))).worksheets[0]

        for column, expected in (
            ("J", dt.date(2020, 3, 13)),
            ("M", dt.date(1990, 5, 17)),
            ("AD", dt.date(2015, 6, 30)),
        ):
            cell = sheet[f"{column}2"]
            assert cell.is_date, column
            assert cell.value.date() == expected, column
            assert cell.number_format == "DD.MM.YYYY", column

    def test_empty_values_stay_empty_cells(self) -> None:
        row = _row(last_name="Иванов", first_name="Иван", phone=79990234365, email="i@mail.ru")
        sheet = _open(render_lms_users_xlsx(_dataset(row))).worksheets[0]

        filled = {cell.column_letter for cell in sheet[2] if cell.value is not None}
        assert filled == {"A", "B", "D", "E"}
        assert sheet["C2"].value is None
        assert sheet["L2"].value is None

    def test_rows_keep_their_order_and_count(self) -> None:
        rows = [_row(last_name=f"Фамилия{n}", first_name="Имя") for n in range(5)]
        sheet = _open(render_lms_users_xlsx(_dataset(*rows))).worksheets[0]

        assert sheet.max_row == 6
        assert [sheet[f"A{n}"].value for n in range(2, 7)] == [f"Фамилия{n}" for n in range(5)]

    def test_there_are_no_formulas_anywhere(self) -> None:
        row = _row(
            last_name="=cmd|'/c calc'!A1",
            first_name="+1+1",
            middle_name="-1-1",
            email="@SUM(A1)",
            reg_street='=HYPERLINK("http://evil")',
            diploma_institution="Обычное название",
        )
        workbook = _open(render_lms_users_xlsx(_dataset(row)))

        for sheet in workbook.worksheets:
            for line in sheet.iter_rows():
                for cell in line:
                    assert cell.data_type != "f", cell.coordinate
        users = workbook["Лист1"]
        assert users["A2"].value == "'=cmd|'/c calc'!A1"
        assert users["B2"].value == "'+1+1"
        assert users["C2"].value == "'-1-1"
        assert users["E2"].value == "'@SUM(A1)"
        assert users["P2"].value == '\'=HYPERLINK("http://evil")'
        assert users["Y2"].value == "Обычное название"

    def test_second_sheet_holds_the_dropdown_lookups(self) -> None:
        sheet = _open(render_lms_users_xlsx(_dataset())).worksheets[1]

        assert [sheet[f"A{n}"].value for n in range(1, 4)] == ["М", "Ж", None]
        assert [sheet[f"B{n}"].value for n in range(1, 8)] == [
            "Без образования",
            "Основное общее образование - 9 классов",
            "Среднее общее образование - 11 классов",
            "Среднее профессиональное образование",
            "Высшее образование – бакалавриат",
            "Высшее образование – специалитет, магистратура",
            "Высшее образование – подготовка кадров высшей квалификации",
        ]
        assert sheet.max_row == 7

    def test_sex_and_education_have_dropdown_validations(self) -> None:
        sheet = _open(render_lms_users_xlsx(_dataset(self._full_row()))).worksheets[0]

        found = {
            str(validation.sqref): validation
            for validation in sheet.data_validations.dataValidation
        }
        assert set(found) == {"L2:L1001", "W2:W1001"}
        sex, education = found["L2:L1001"], found["W2:W1001"]
        assert (sex.type, sex.formula1) == ("list", "Лист2!$A$1:$A$2")
        assert (education.type, education.formula1) == ("list", "Лист2!$B$1:$B$7")
        for validation in (sex, education):
            assert validation.allow_blank is True
            assert validation.showErrorMessage is True

    def test_render_report_dispatches_to_the_lms_template_only_for_its_kind(self) -> None:
        content = render_report(_dataset(), format="xlsx", template_code="lms_users_upload")
        assert _open(content).sheetnames == ["Лист1", "Лист2"]

        # Прочие отчёты по-прежнему пишутся общим рендером: один лист с названием отчёта.
        generic = render_report(_dataset(), format="xlsx", template_code="kam_summary")
        assert _open(generic).sheetnames == ["Загрузка пользователей"]


# =============================================================================
# Выборка и путь через очередь отчётов (настоящая Postgres)
# =============================================================================


def _marker() -> str:
    return uuid.uuid4().hex[:8]


async def _seed_world(head_id: uuid.UUID, other_id: uuid.UUID) -> dict[str, Any]:
    """Воронка физлиц и люди с разными сделками. Ключи результата — роли людей в тесте."""
    from app.core.db import session_scope
    from app.modules.catalog.models import (
        Contact,
        ContactLearnerProfile,
        Organization,
        Product,
    )
    from app.modules.crm.models import Deal, DealProduct
    from app.modules.workflow.models import Workflow, WorkflowStatus

    mark = _marker()
    async with session_scope() as session:
        workflow = Workflow(code=f"b2c-{mark}", name="Воронка физлиц", deal_type="b2c")
        session.add(workflow)
        await session.flush()
        statuses = {}
        for order, (code, name) in enumerate(
            (
                ("site_application", "Заявка с сайта"),
                ("payment_contract", "Оплата и договор оферты"),
                ("lms_enrollment", "Зачисление в LMS"),
            )
        ):
            statuses[code] = WorkflowStatus(
                workflow_id=workflow.id, code=code, name=name, sort_order=order
            )
            session.add(statuses[code])
        product_a = Product(code=f"course-a-{mark}", name=f"Курс А {mark}")
        product_b = Product(code=f"course-b-{mark}", name=f"Курс Б {mark}")
        organization = Organization(name=f"Вуз {mark}", org_type="university")
        session.add_all([product_a, product_b, organization])
        await session.flush()

        def contact(key: str, **fields: Any) -> Contact:
            person = Contact(
                first_name="Иван",
                last_name=f"{key}{mark}",
                email=f"{key.lower()}.{mark}@example.ru",
                phone=f"+7999{int(uuid.uuid4().hex[:8], 16) % 10_000_000:07d}",
                **fields,
            )
            session.add(person)
            return person

        people = {
            "Аня": contact("Аня"),  # оплата, курс А поток 1, полный профиль
            "Борис": contact("Борис"),  # зачисление, курс А поток 2, профиля нет
            "Вера": contact("Вера"),  # только заявка с сайта: этап не оплачен
            "Глеб": contact("Глеб"),  # оплата, но сделка чужого менеджера
            "Дарья": contact("Дарья"),  # две оплаченные сделки: в файле одна строка
            "Егор": contact("Егор"),  # оплата, но контакт удалён
            "Жанна": contact("Жанна"),  # сделка B2B: не физлицо
        }
        people["Егор"].deleted_at = dt.datetime.now(dt.UTC)
        await session.flush()

        session.add(
            ContactLearnerProfile(
                contact_id=people["Аня"].id,
                snils="11223344595",
                passport_series="4512",
                passport_number="123456",
                passport_issued_by="ОУФМС России по г. Москве",
                passport_issued_at=dt.date(2020, 3, 13),
                passport_dept_code="770-001",
                sex="F",
                birth_date=dt.date(1990, 5, 17),
                reg_region="Москва",
                reg_city="Москва",
                reg_street="Тверская",
                reg_house="1",
                reg_apartment="15",
                reg_zip="125009",
                first_name_dative="Ане",
                education="higher_bachelor",
                diploma_number="1234567",
                diploma_issued_at=dt.date(2015, 6, 30),
            )
        )

        rows: list[tuple[Deal, Product | None, int | None]] = []

        def deal(
            person: Contact,
            status: str,
            owner_id: uuid.UUID,
            *,
            product: Product | None = None,
            stream: int | None = None,
            deal_type: str = "b2c",
            created_at: dt.datetime | None = None,
        ) -> Deal:
            row = Deal(
                number=f"D-{uuid.uuid4().hex[:12]}",
                title="Оплаченный курс",
                deal_type=deal_type,
                workflow_id=workflow.id,
                status_id=statuses[status].id,
                contact_id=person.id,
                organization_id=organization.id if deal_type == "b2b" else None,
                owner_id=owner_id,
                **({"created_at": created_at} if created_at else {}),
            )
            session.add(row)
            rows.append((row, product, stream))
            return row

        deal(people["Аня"], "payment_contract", head_id, product=product_a, stream=1)
        deal(people["Борис"], "lms_enrollment", head_id, product=product_a, stream=2)
        deal(people["Вера"], "site_application", head_id, product=product_a, stream=1)
        deal(people["Глеб"], "payment_contract", other_id, product=product_a, stream=1)
        deal(
            people["Дарья"],
            "payment_contract",
            head_id,
            product=product_b,
            stream=1,
            created_at=dt.datetime(2025, 1, 10, 12, 0, tzinfo=dt.UTC),
        )
        deal(
            people["Дарья"],
            "lms_enrollment",
            head_id,
            product=product_b,
            stream=3,
            created_at=dt.datetime(2025, 2, 20, 12, 0, tzinfo=dt.UTC),
        )
        deal(people["Егор"], "payment_contract", head_id, product=product_a, stream=1)
        deal(people["Жанна"], "payment_contract", head_id, deal_type="b2b")
        await session.flush()
        for row, product, stream in rows:
            if product is not None:
                session.add(
                    DealProduct(
                        deal_id=row.id, product_id=product.id, quantity=1, stream_number=stream
                    )
                )
        await session.flush()
        return {
            "mark": mark,
            "names": {key: person.last_name for key, person in people.items()},
            "product_a": product_a.id,
            "product_b": product_b.id,
            "anya_phone": people["Аня"].phone,
            "anya_email": people["Аня"].email,
        }


async def _build(user_id: uuid.UUID, params: dict[str, Any]) -> ReportDataset:
    from app.core.db import session_scope
    from app.modules.identity.models import User
    from app.modules.reporting.builders import build_lms_users_upload
    from app.modules.reporting.tasks import _principal_for_worker

    async with session_scope() as session:
        user = await session.get(User, user_id)
        assert user is not None
        return await build_lms_users_upload(session, _principal_for_worker(user), params)


@pytest.fixture
def world(client) -> dict[str, Any]:
    """Руководитель (принципал отчёта), чужой менеджер и данные из `_seed_world`."""
    head = run(client, _make_user, "HEAD")
    other = run(client, _make_user, "KAM")
    data = run(client, _seed_world, head.id, other.id)
    data["head"], data["other"] = head, other
    return data


def _names(dataset: ReportDataset, world: dict[str, Any]) -> set[str]:
    """Чьи строки попали в файл, по ключам людей из `_seed_world`."""
    by_last_name = {last_name: key for key, last_name in world["names"].items()}
    return {by_last_name[row[0]] for row in dataset.rows}


@needs_db
class TestSelection:
    def test_rows_are_the_paid_b2c_contacts_of_the_principals_deals(self, client, world) -> None:
        dataset = run(client, _build, world["head"].id, {})

        # Оплаченные и зачисленные физлица руководителя — по одной строке на человека
        # (у Дарьи две подходящие сделки). Остальные отсеяны: этап, чужая сделка, удалённый
        # контакт, B2B.
        assert _names(dataset, world) == {"Аня", "Борис", "Дарья"}
        assert len(dataset.rows) == 3
        assert dataset.columns == HEADERS

    def test_rows_are_sorted_by_surname_and_carry_the_template_columns(self, client, world) -> None:
        dataset = run(client, _build, world["head"].id, {})

        assert [row[0] for row in dataset.rows] == sorted(row[0] for row in dataset.rows)
        assert all(len(row) == 30 for row in dataset.rows)

    def test_a_contact_from_a_deal_outside_the_scope_is_excluded(self, client, world) -> None:
        # Глеба оплатил менеджер вне команды руководителя: в его отчёт он не попадает.
        head_rows = run(client, _build, world["head"].id, {})
        assert "Глеб" not in _names(head_rows, world)

        # Свой отчёт у чужого менеджера — только его собственные сделки.
        other_rows = run(client, _build, world["other"].id, {})
        assert _names(other_rows, world) == {"Глеб"}

    def test_profile_values_are_written_in_template_form(self, client, world) -> None:
        dataset = run(client, _build, world["head"].id, {})
        row = next(row for row in dataset.rows if row[0] == world["names"]["Аня"])
        cells = dict(zip(TARGETS, row, strict=True))

        assert cells["phone"] == int(world["anya_phone"].lstrip("+"))
        assert cells["email"] == world["anya_email"]
        assert cells["snils"] == "112-233-445 95"
        assert cells["passport_series"] == "4512"
        assert cells["passport_dept_code"] == "770-001"
        assert cells["passport_issued_at"] == dt.date(2020, 3, 13)
        assert cells["sex"] == "Ж"
        assert cells["birth_date"] == dt.date(1990, 5, 17)
        assert cells["education"] == "Высшее образование – бакалавриат"
        assert cells["diploma_issued_at"] == dt.date(2015, 6, 30)
        assert cells["middle_name"] is None

    def test_contact_without_a_profile_has_only_the_basic_columns(self, client, world) -> None:
        dataset = run(client, _build, world["head"].id, {})
        row = next(row for row in dataset.rows if row[0] == world["names"]["Борис"])
        cells = dict(zip(TARGETS, row, strict=True))

        assert {target for target, value in cells.items() if value is not None} == {
            "last_name",
            "first_name",
            "phone",
            "email",
        }

    @pytest.mark.parametrize(
        ("params", "expected"),
        [
            ({"product_id": "A"}, {"Аня", "Борис"}),
            ({"product_id": "B"}, {"Дарья"}),
            ({"product_id": "A", "stream_number": 2}, {"Борис"}),
            ({"stream_number": 1}, {"Аня", "Дарья"}),
            # Продукт и поток — об одной позиции сделки: у Дарьи поток 3 у курса Б, а не у А.
            ({"product_id": "A", "stream_number": 3}, set()),
            ({"product_id": "B", "stream_number": 3}, {"Дарья"}),
            ({"status_codes": ["site_application"]}, {"Вера"}),
            ({"status_codes": ["lms_enrollment"]}, {"Борис", "Дарья"}),
            ({"status_codes": "payment_contract, lms_enrollment"}, {"Аня", "Борис", "Дарья"}),
            ({"date_from": "2025-01-01", "date_to": "2025-01-31"}, {"Дарья"}),
            ({"date_from": "2025-02-01", "date_to": "2025-02-28"}, {"Дарья"}),
            ({"date_to": "2025-12-31"}, {"Дарья"}),
            ({"date_from": "2025-01-11", "date_to": "2025-02-19"}, set()),
        ],
    )
    def test_filters(self, client, world, params, expected) -> None:
        params = dict(params)
        if params.get("product_id") in ("A", "B"):
            params["product_id"] = str(world[f"product_{params['product_id'].lower()}"])

        dataset = run(client, _build, world["head"].id, params)

        assert _names(dataset, world) == expected

    def test_date_filter_is_inclusive_on_both_ends(self, client, world) -> None:
        dataset = run(
            client, _build, world["head"].id, {"date_from": "2025-01-10", "date_to": "2025-01-10"}
        )
        assert _names(dataset, world) == {"Дарья"}

    @pytest.mark.parametrize(
        "params",
        [
            {"stream_number": 0},
            {"stream_number": -5},
            {"stream_number": "первый"},
            {"stream_number": True},
            {"stream_number": 10**12},
            {"product_id": "не-uuid"},
            {"status_codes": 5},
            {"status_codes": [1, 2]},
            {"status_codes": ["x" * 65]},
            {"date_from": "2025-02-01", "date_to": "2025-01-01"},
            {"date_from": "вчера"},
        ],
    )
    def test_invalid_params_are_refused(self, client, world, params) -> None:
        from app.core.errors import ValidationError

        with pytest.raises(ValidationError):
            run(client, _build, world["head"].id, params)


@needs_db
class TestThroughTheReportQueue:
    """`POST /api/reports` → файл в хранилище → запись аудита `REPORT_EXPORTED`."""

    def _stub_storage(self, monkeypatch) -> list[bytes]:
        import app.modules.reporting.service as reporting_service

        uploaded: list[bytes] = []

        async def fake_ensure_bucket(bucket: str) -> None:
            return None

        async def fake_upload(*, bucket: str, key: str, body: bytes, content_type: str) -> None:
            uploaded.append(body)

        monkeypatch.setattr(reporting_service, "ensure_bucket", fake_ensure_bucket)
        monkeypatch.setattr(reporting_service, "upload_object_bytes", fake_upload)
        return uploaded

    def _seed_templates(self, client) -> None:
        from app.core.db import session_scope
        from app.modules.reporting.seed import seed_report_templates

        async def _seed() -> None:
            async with session_scope() as session:
                await seed_report_templates(session)

        run(client, _seed)

    def _sign_in(self, client, user) -> None:
        client.headers["X-CSRF-Token"] = authenticate(client, user)

    def _create(self, client, **params):
        return client.post(
            "/api/reports",
            json={"template_code": "lms_users_upload", "format": "xlsx", "params": params},
        )

    def test_head_gets_the_template_file_and_the_export_is_audited(
        self, client, world, monkeypatch
    ) -> None:
        from tests.people_helpers import audit_entries

        uploaded = self._stub_storage(monkeypatch)
        self._seed_templates(client)
        self._sign_in(client, world["head"])

        created = self._create(client)
        assert created.status_code == 201, created.text
        job = created.json()
        assert job["status"] == "completed"
        assert job["row_count"] == 3
        assert job["format"] == "xlsx"

        [content] = uploaded
        sheet = _open(content).worksheets[0]
        assert [cell.value for cell in sheet[1]] == HEADERS
        assert sheet.max_row == 4
        anya = next(
            row for row in sheet.iter_rows(min_row=2) if row[0].value == world["names"]["Аня"]
        )
        assert anya[3].value == int(world["anya_phone"].lstrip("+"))
        assert anya[5].value == "112-233-445 95"
        assert anya[11].value == "Ж"
        assert anya[12].is_date
        assert anya[12].number_format == "DD.MM.YYYY"

        [entry] = audit_entries(client, "REPORT_EXPORTED", job["id"])
        assert entry["changes"]["template_code"]["new"] == "lms_users_upload"
        assert entry["changes"]["row_count"]["new"] == 3
        assert any("ПДн" in category for category in entry["changes"]["data_categories"]["new"])

    def test_params_narrow_the_file(self, client, world, monkeypatch) -> None:
        uploaded = self._stub_storage(monkeypatch)
        self._seed_templates(client)
        self._sign_in(client, world["head"])

        created = self._create(
            client,
            product_id=str(world["product_a"]),
            stream_number=2,
            status_codes=["lms_enrollment"],
        )
        assert created.status_code == 201, created.text
        assert created.json()["row_count"] == 1
        sheet = _open(uploaded[-1]).worksheets[0]
        assert sheet["A2"].value == world["names"]["Борис"]

    def test_bad_params_are_a_422_and_leave_no_job(self, client, world, monkeypatch) -> None:
        self._stub_storage(monkeypatch)
        self._seed_templates(client)
        self._sign_in(client, world["head"])

        response = self._create(client, stream_number=0)
        assert response.status_code == 422, response.text
        assert response.json()["errors"][0]["field"] == "stream_number"
        assert (
            client.get("/api/reports", params={"template_code": "lms_users_upload"}).json()["items"]
            == []
        )

    def test_only_xlsx_is_offered(self, client, world, monkeypatch) -> None:
        self._stub_storage(monkeypatch)
        self._seed_templates(client)
        self._sign_in(client, world["head"])

        for fmt in ("pdf", "png"):
            response = client.post(
                "/api/reports", json={"template_code": "lms_users_upload", "format": fmt}
            )
            assert response.status_code == 422, (fmt, response.text)

    def test_kam_cannot_run_it_and_does_not_see_the_template(
        self, client, world, monkeypatch
    ) -> None:
        self._stub_storage(monkeypatch)
        self._seed_templates(client)
        self._sign_in(client, world["other"])  # менеджер (KAM)

        response = self._create(client)
        assert response.status_code == 403, response.text
        assert response.json()["code"] == "CRM-1102"
        listed = client.get("/api/report-templates").json()["items"]
        assert "lms_users_upload" not in {item["code"] for item in listed}

        self._sign_in(client, world["head"])
        listed = client.get("/api/report-templates").json()["items"]
        assert "lms_users_upload" in {item["code"] for item in listed}

    def test_admin_may_run_it(self, client, world, monkeypatch) -> None:
        self._stub_storage(monkeypatch)
        self._seed_templates(client)
        admin = run(client, _make_user, "ADMIN")
        self._sign_in(client, admin)

        # Без фильтра админ видит сделки всех; ограничиваем продуктом теста.
        created = self._create(client, product_id=str(world["product_a"]))
        assert created.status_code == 201, created.text
        assert created.json()["row_count"] == 3  # Аня, Борис и Глеб: у админа скоуп «все»

    def test_json_data_of_the_pii_report_is_refused(self, client, world, monkeypatch) -> None:
        self._stub_storage(monkeypatch)
        self._seed_templates(client)
        self._sign_in(client, world["head"])
        job_id = self._create(client).json()["id"]

        # У `/data` нет записи `REPORT_EXPORTED`: ПДн целиком отдаются только файлом.
        response = client.get(f"/api/reports/{job_id}/data")
        assert response.status_code == 403, response.text
