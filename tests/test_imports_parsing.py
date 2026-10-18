"""B3-2: листы книги и строка заголовков (без БД — книги собирает openpyxl в памяти).

Реальный файл заказчика — книга из нескольких листов с титульной строкой
перед шапкой, иногда со скрытым листом (черновик/архив). До этой задачи
читался только первый лист и первая строка.

Полный контракт задачи: выбор листа, автопоиск строки заголовка, разбор с
явно указанной строкой заголовка, абсолютная нумерация строк Excel
(`source_rows`), склейка нескольких листов (`sheets=["*"]`) с проверкой
одинаковой шапки и `fill_down` для объединённых ячеек.
"""

from __future__ import annotations

import io

import openpyxl
import pytest
import xlwt

from app.core.errors import AppError
from app.modules.imports.parsing import detect_header_row, list_sheets, parse_table

_LICENSE_LABELS = [
    "Название ВУЗа",
    "Вендор",
    "ПО",
    "Номер договора",
    "Подписание лицензии",
    "Срок действия лицензии (год)",
    "Статус по передаче",
    "ФИО Менеджера",
    "Ответственные от ВУЗа",
    "Комментарий",
]


def _xlsx_bytes(sheets: dict[str, list[list[object]]], *, hidden: set[str] = frozenset()) -> bytes:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
        if name in hidden:
            sheet.sheet_state = "hidden"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _xls_bytes(rows: list[list[object]]) -> bytes:
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Лист1")
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            sheet.write(r, c, value)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class TestDetectHeaderRow:
    def test_header_on_first_row_when_it_matches(self) -> None:
        rows = [["Название ВУЗа", "Вендор", "ПО"], ["МГТУ", "Яндекс", "DataLake"]]
        assert detect_header_row(rows, _LICENSE_LABELS) == 1

    def test_header_after_title_and_blank_rows(self) -> None:
        rows = [
            ["Каталог лицензий 2026"],
            [],
            ["Название ВУЗа", "Вендор", "ПО", "Номер договора"],
            ["МГТУ", "Яндекс", "DataLake", "Д-1"],
        ]
        assert detect_header_row(rows, _LICENSE_LABELS) == 3

    def test_too_few_matches_falls_back_to_first_row(self) -> None:
        rows = [["совсем", "другие", "колонки"], ["x", "y", "z"]]
        assert detect_header_row(rows, _LICENSE_LABELS) == 1

    def test_tolerates_a_typo_in_one_header(self) -> None:
        rows = [["Названи ВУЗа", "Вендор", "ПО", "Номер договора"], ["МГТУ", "Яндекс", "DL", "Д-1"]]
        assert detect_header_row(rows, _LICENSE_LABELS) == 1

    def test_no_expected_labels_defaults_to_first_row(self) -> None:
        assert detect_header_row([["a"], ["b"]], []) == 1


class TestListSheetsXlsx:
    def test_names_and_row_counts(self) -> None:
        content = _xlsx_bytes(
            {
                "Лист1": [["A", "B"], ["1", "2"]],
                "Лист2": [["C"], ["3"], ["4"], ["5"]],
            }
        )
        sheets = list_sheets(content, source_format="xlsx")
        by_name = {s.name: s for s in sheets}
        assert by_name["Лист1"].rows == 2
        assert by_name["Лист2"].rows == 4
        assert by_name["Лист1"].visible
        assert by_name["Лист2"].visible

    def test_hidden_sheet_is_marked_not_visible(self) -> None:
        content = _xlsx_bytes(
            {"Видимый": [["A"], ["1"]], "Архив": [["A"], ["2"]]}, hidden={"Архив"}
        )
        sheets = {s.name: s for s in list_sheets(content, source_format="xlsx")}
        assert sheets["Видимый"].visible
        assert not sheets["Архив"].visible

    def test_suggested_header_row_uses_expected_labels(self) -> None:
        content = _xlsx_bytes(
            {
                "Лицензии": [
                    ["Каталог на 2026 год"],
                    ["Название ВУЗа", "Вендор", "ПО", "Номер договора"],
                    ["МГТУ", "Яндекс", "DataLake", "Д-1"],
                ]
            }
        )
        sheets = list_sheets(content, source_format="xlsx", expected_labels=_LICENSE_LABELS)
        assert sheets[0].suggested_header_row == 2


class TestParseTableWithSheetAndHeaderRow:
    def test_default_behaviour_is_unchanged(self) -> None:
        content = _xlsx_bytes({"Лист1": [["Название ВУЗа", "Вендор"], ["МГТУ", "Яндекс"]]})
        table = parse_table(content, source_format="xlsx")
        assert table.headers == ["Название ВУЗа", "Вендор"]
        assert table.rows == [["МГТУ", "Яндекс"]]

    def test_reads_a_named_sheet(self) -> None:
        content = _xlsx_bytes(
            {
                "Лист1": [["x"], ["1"]],
                "Лист2": [["Название ВУЗа", "Вендор"], ["СПбГУ", "Ростелеком"]],
            }
        )
        table = parse_table(content, source_format="xlsx", sheet="Лист2")
        assert table.headers == ["Название ВУЗа", "Вендор"]
        assert table.rows == [["СПбГУ", "Ростелеком"]]
        assert table.sheet_name == "Лист2"

    def test_reads_header_from_an_explicit_row(self) -> None:
        content = _xlsx_bytes(
            {
                "Лицензии": [
                    ["Каталог на 2026 год"],
                    ["Название ВУЗа", "Вендор"],
                    ["МГТУ", "Яндекс"],
                    ["СПбГУ", "Ростелеком"],
                ]
            }
        )
        table = parse_table(content, source_format="xlsx", header_row=2)
        assert table.headers == ["Название ВУЗа", "Вендор"]
        assert table.rows == [["МГТУ", "Яндекс"], ["СПбГУ", "Ростелеком"]]
        assert table.header_row == 2

    def test_unknown_sheet_name_is_a_clear_error(self) -> None:
        content = _xlsx_bytes({"Лист1": [["a"], ["1"]]})
        with pytest.raises(AppError):
            parse_table(content, source_format="xlsx", sheet="Нет такого")

    def test_header_row_out_of_range_is_a_clear_error(self) -> None:
        content = _xlsx_bytes({"Лист1": [["a"], ["1"]]})
        with pytest.raises(AppError):
            parse_table(content, source_format="xlsx", header_row=50)


class TestSourceRows:
    """B3-2: `source_rows`/`sheet_names` — абсолютный номер строки Excel и имя листа на каждую
    строку результата, для отчёта об ошибках (B3-5)."""

    def test_header_after_title_and_blank_rows_gives_absolute_source_row(self) -> None:
        content = _xlsx_bytes(
            {
                "Лицензии": [
                    ["Каталог лицензий 2026"],
                    [],
                    [],
                    ["Название ВУЗа", "Вендор"],
                    ["МГТУ", "Яндекс"],
                ]
            }
        )
        table = parse_table(content, source_format="xlsx", header_row=4)
        assert table.rows == [["МГТУ", "Яндекс"]]
        assert table.source_rows == [5]
        assert table.sheet_names == ["Лицензии"]

    def test_blank_row_inside_data_keeps_the_gap_in_source_rows(self) -> None:
        content = _xlsx_bytes(
            {
                "Лист1": [
                    ["Название ВУЗа", "Вендор"],
                    ["МГТУ", "Яндекс"],
                    [],
                    ["СПбГУ", "Ростелеком"],
                ]
            }
        )
        table = parse_table(content, source_format="xlsx")
        assert table.rows == [["МГТУ", "Яндекс"], ["СПбГУ", "Ростелеком"]]
        # Строка 2 — первая строка данных, строка 3 пуста и пропущена, строка 4 — вторая.
        assert table.source_rows == [2, 4]

    def test_default_single_sheet_name_is_populated(self) -> None:
        content = _xlsx_bytes({"Лист1": [["a"], ["1"]]})
        table = parse_table(content, source_format="xlsx")
        assert table.sheet_names == ["Лист1"]


class TestFillDown:
    """B3-2: объединённые ячейки (вуз/вендор записаны один раз на блок строк)."""

    def test_fills_empty_cell_from_the_row_above(self) -> None:
        content = _xlsx_bytes(
            {
                "Лист1": [
                    ["Название ВУЗа", "ПО"],
                    ["МГТУ", "DataLake"],
                    ["", "CRM"],
                    ["СПбГУ", "ERP"],
                ]
            }
        )
        table = parse_table(content, source_format="xlsx", fill_down=["Название ВУЗа"])
        assert table.rows == [["МГТУ", "DataLake"], ["МГТУ", "CRM"], ["СПбГУ", "ERP"]]
        assert table.filled_cells == [(3, "Название ВУЗа")]

    def test_fully_blank_row_resets_the_fill(self) -> None:
        content = _xlsx_bytes(
            {
                "Лист1": [
                    ["Название ВУЗа", "ПО"],
                    ["МГТУ", "DataLake"],
                    [],
                    ["", "ERP"],
                ]
            }
        )
        table = parse_table(content, source_format="xlsx", fill_down=["Название ВУЗа"])
        # Пустая строка (строка 3) разделяет блоки — колонка в строке 4 остаётся пустой.
        assert table.rows == [["МГТУ", "DataLake"], ["", "ERP"]]
        assert table.filled_cells == []

    def test_fill_down_works_on_xls_too(self) -> None:
        content = _xls_bytes([["Название ВУЗа", "ПО"], ["МГТУ", "DataLake"], ["", "CRM"]])
        table = parse_table(content, source_format="xls", fill_down=["Название ВУЗа"])
        assert table.rows == [["МГТУ", "DataLake"], ["МГТУ", "CRM"]]


class TestMergeAllSheets:
    """B3-2: `sheets=["*"]` склеивает все видимые листы с одинаковой шапкой."""

    def test_merges_two_sheets_with_the_same_header(self) -> None:
        content = _xlsx_bytes(
            {
                "2024": [["Название ВУЗа", "ПО"], ["МГТУ", "DataLake"]],
                "2025": [["Название ВУЗа", "ПО"], ["СПбГУ", "ERP"]],
            }
        )
        table = parse_table(content, source_format="xlsx", sheets=["*"])
        assert table.headers == ["Название ВУЗа", "ПО"]
        assert table.rows == [["МГТУ", "DataLake"], ["СПбГУ", "ERP"]]
        assert table.sheet_names == ["2024", "2025"]
        assert table.source_rows == [2, 2]

    def test_hidden_sheet_is_excluded_from_merge(self) -> None:
        content = _xlsx_bytes(
            {
                "Видимый": [["Название ВУЗа"], ["МГТУ"]],
                "Архив": [["Название ВУЗа"], ["Черновик"]],
            },
            hidden={"Архив"},
        )
        table = parse_table(content, source_format="xlsx", sheets=["*"])
        assert table.rows == [["МГТУ"]]

    def test_mismatched_headers_raise_a_clear_error(self) -> None:
        content = _xlsx_bytes(
            {
                "Лист1": [["Название ВУЗа", "ПО"], ["МГТУ", "DataLake"]],
                "Лист2": [["Вуз", "Продукт"], ["СПбГУ", "ERP"]],
            }
        )
        with pytest.raises(AppError):
            parse_table(content, source_format="xlsx", sheets=["*"])

    def test_explicit_header_row_applies_to_every_sheet(self) -> None:
        content = _xlsx_bytes(
            {
                "A": [["Титул"], ["Название ВУЗа"], ["МГТУ"]],
                "B": [["Титул"], ["Название ВУЗа"], ["СПбГУ"]],
            }
        )
        table = parse_table(content, source_format="xlsx", sheets=["*"], header_row=2)
        assert table.rows == [["МГТУ"], ["СПбГУ"]]
        assert table.source_rows == [3, 3]


class TestImportOptions:
    def test_sheet_and_sheets_are_mutually_exclusive(self) -> None:
        import pytest
        from pydantic import ValidationError

        from app.modules.imports.schemas import ImportOptions

        with pytest.raises(ValidationError, match="взаимоисключающие"):
            ImportOptions(sheet="2024", sheets=["*"])

        assert ImportOptions(sheet="2024").sheets is None
        assert ImportOptions(sheets=["*"]).sheet is None
        assert ImportOptions(sheet="2024", sheets=[]).sheets == []
