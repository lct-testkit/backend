"""B3-9: шаблоны для скачивания (xlsx/xls) и встроенные пресеты маппинга."""

from __future__ import annotations

import io

import openpyxl
import pytest
import xlrd

from app.modules.imports.fields import LICENSE_FIELDS
from tests.conftest import TEST_DATABASE_URL, run
from tests.crm_helpers import login
from tests.import_helpers import job_rows, stub_storage, token

pytestmark = pytest.mark.skipif(
    not TEST_DATABASE_URL, reason="нужен TEST_DATABASE_URL с применёнными миграциями"
)


class TestDownloadTemplate:
    def test_xlsx_template_has_license_headers_in_order(self, client) -> None:
        login(client, "HEAD")
        response = client.get("/api/imports/templates/license", params={"format": "xlsx"})
        assert response.status_code == 200, response.text
        assert response.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        assert "Данные" in workbook.sheetnames
        assert "Инструкция" in workbook.sheetnames
        sheet = workbook["Данные"]
        headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
        assert headers == [f.label for f in LICENSE_FIELDS]
        # Хотя бы одна учебная строка.
        assert sheet.max_row >= 2

    def test_xls_template_opens_with_xlrd(self, client) -> None:
        login(client, "HEAD")
        response = client.get("/api/imports/templates/license", params={"format": "xls"})
        assert response.status_code == 200, response.text
        assert response.content[:4] == b"\xd0\xcf\x11\xe0"  # OLE/CFB — "Composite Document File"
        book = xlrd.open_workbook(file_contents=response.content)
        assert "Данные" in book.sheet_names()
        sheet = book.sheet_by_name("Данные")
        headers = [sheet.cell_value(0, i) for i in range(sheet.ncols)]
        assert headers == [f.label for f in LICENSE_FIELDS]

    def test_direction_and_responsible_templates_download_too(self, client) -> None:
        login(client, "ADMIN")
        for entity_type in ("direction", "responsible"):
            response = client.get(f"/api/imports/templates/{entity_type}")
            assert response.status_code == 200, response.text

    def test_kam_is_forbidden(self, client) -> None:
        login(client, "KAM")
        response = client.get("/api/imports/templates/license")
        assert response.status_code == 403, response.text

    def test_unknown_entity_type_is_422(self, client) -> None:
        login(client, "HEAD")
        response = client.get("/api/imports/templates/not-a-type")
        assert response.status_code == 422, response.text

    def test_unknown_format_is_422(self, client) -> None:
        login(client, "HEAD")
        response = client.get("/api/imports/templates/license", params={"format": "pdf"})
        assert response.status_code == 422, response.text


class TestTemplateRoundTrip:
    def test_downloaded_xlsx_template_maps_all_columns_and_dry_runs_clean(
        self, client, monkeypatch
    ) -> None:
        """«Скачать -> дополнить -> dry-run» (B3-9 приёмка): шаблон уже содержит две годные
        учебные строки, поэтому достаточно скачать его и загрузить обратно без правок."""
        from tests.import_helpers import _seed_file

        tok = token()
        admin = login(client, "ADMIN")

        async def seed_university(name: str):
            from app.core.db import session_scope
            from app.modules.catalog.models import Organization

            async with session_scope() as session:
                session.add(Organization(name=name, org_type="university"))
                await session.flush()

        # Шаблон называет вуз «МГТУ им. Н.Э. Баумана»/«СПбГУ» — заводим их, чтобы org_not_found
        # не испортил чистый dry-run (имена шаблона фиксированы, не параметризуются `tok`).
        run(client, seed_university, "МГТУ им. Н.Э. Баумана")
        run(client, seed_university, "СПбГУ")

        template = client.get("/api/imports/templates/license", params={"format": "xlsx"})
        assert template.status_code == 200, template.text
        content = template.content

        stub_storage(monkeypatch, content)
        file_id = run(
            client,
            _seed_file,
            content,
            admin.id,
            f"template-{tok}.xlsx",
            "application/octet-stream",
        )
        created = client.post(
            "/api/imports",
            json={
                "file_id": str(file_id),
                "entity_type": "license",
                "mode": "upsert",
                "source_format": "xlsx",
            },
        )
        assert created.status_code == 201, created.text
        job_id = created.json()["id"]

        profile = client.get(f"/api/imports/{job_id}/profile")
        assert profile.status_code == 200, profile.text
        suggested = profile.json()["suggested_mapping"]
        assert len(suggested) == len(LICENSE_FIELDS)

        saved = client.put(f"/api/imports/{job_id}/mapping", json={"mapping": suggested})
        assert saved.status_code == 200, saved.text
        dry = client.post(f"/api/imports/{job_id}/dry-run")
        assert dry.status_code == 200, dry.text
        assert dry.json()["error_rows"] == 0, job_rows(client, job_id, status="error")


class TestBuiltinPresets:
    def test_builtin_license_preset_is_listed_without_id(self, client) -> None:
        login(client, "HEAD")
        response = client.get("/api/import-presets", params={"entity_type": "license"})
        assert response.status_code == 200, response.text
        items = response.json()["items"]
        builtin = [p for p in items if p["builtin"]]
        assert any(p["name"] == "Лицензии: 10 колонок ТЗ" for p in builtin)
        assert all(p["id"] is None for p in builtin)
