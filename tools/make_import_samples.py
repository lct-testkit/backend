"""B3-11: образцы файлов импорта для демонстрации и сквозного теста.

    python tools/make_import_samples.py

Пишет в `docs/samples/imports/` (сид `.mjs` для вузов/вендоров и каталоги для деплоя — в
репозиториях frontend/deploy, см. отчёт по эпику B3-11; здесь только backend-часть: файлы,
которые можно скачать со страницы импорта и загрузить на демо-стенде или в тесте).

Все данные выдуманные (152-ФЗ): вузы, вендоры и ФИО — не настоящие; email — в домене `example.
com` (`example.test` синтаксически проходит как email, но реальный `email-validator`,
используемый `imports.fields.validate_field`, отклоняет зарезервированные тестовые TLD — поэтому
не `example.test`, как по умолчанию предлагает документ задачи, а `example.com`, как и остальные
тестовые e-mail в этом репозитории).

Файлы:

* `licenses-10-columns.xlsx`/`.xls` — 30 строк, русские статусы, «да»/«нет», три намеренные
  ошибки (см. `_LICENSE_ROWS`, последние три строки). `tests/test_imports_samples.py` сеет
  названные здесь вузы и проверяет `error_rows == 3`, `ok_rows + warn_rows == 27`.
* `licenses-multisheet.xlsx` — три листа (по годам), шапка в строке 3, на первом листе вендор
  записан в объединённых ячейках (реальный файл заказчика часто экономит так на повторе).
* `directions.xlsx` — направления с одним уровнем вложенности.
* `responsible.xlsx` — список ответственных по email и по ФИО.
"""

from __future__ import annotations

import io
from pathlib import Path

import openpyxl
import xlwt
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parent.parent
SAMPLES_DIR = ROOT / "docs" / "samples" / "imports"

LICENSE_HEADERS = [
    "Название ВУЗа",
    "Вендор",
    "ПО",
    "Номер договора",
    "Подписание лицензии",
    "Срок действия лицензии (год)",
    "Статус по передаче",
    "ФИО Менеджера",
    "Ответственные от ВУЗа",
    "Комментарий",
]

# Пять выдуманных вузов — тест сеет их под этими именами (`tests/test_imports_samples.py`).
UNIVERSITIES = [
    "Московский технический университет «Прогресс»",
    "Санкт-Петербургский политехнический институт «Нева»",
    "Уральский государственный университет «Исеть»",
    "Сибирский технологический университет «Енисей»",
    "Южный федеральный университет «Дон»",
]
VENDORS = ["ООО «Базис Технологии»", "ПАО «Инфосвязь»", "ООО «Датапоток»"]
PRODUCTS = ["ИС.Документооборот", "ИС.Аналитика", "ИС.Хранилище"]
STATUSES = ["передано", "в работе", "не начато", "передана"]
MANAGERS = ["Иванов Иван Иванович", "Петрова Анна Сергеевна", "Сидоров Пётр Николаевич"]


def _good_license_rows(count: int) -> list[list[str]]:
    rows = []
    for i in range(count):
        university = UNIVERSITIES[i % len(UNIVERSITIES)]
        vendor = VENDORS[i % len(VENDORS)]
        product = PRODUCTS[i % len(PRODUCTS)]
        status = STATUSES[i % len(STATUSES)]
        manager = MANAGERS[i % len(MANAGERS)]
        signed = "да" if i % 2 == 0 else "15.03.2024"
        term = str(1 + (i % 5)) if i % 3 else "бессрочно"
        rows.append(
            [
                university,
                vendor,
                product,
                f"Д-{100 + i}/2026",
                signed,
                term,
                status,
                manager,
                "",
                "" if i % 4 else "Образец для демонстрации",
            ]
        )
    return rows


def _bad_license_rows() -> list[list[str]]:
    return [
        # 1. Пустое название вуза -> required_empty.
        ["", VENDORS[0], PRODUCTS[0], "Д-900/2026", "да", "3", "передано", MANAGERS[0], "", ""],
        # 2. Неизвестный статус по передаче -> status_unknown.
        [
            UNIVERSITIES[0],
            VENDORS[1],
            PRODUCTS[1],
            "Д-901/2026",
            "да",
            "3",
            "абракадабра",
            MANAGERS[1],
            "",
            "",
        ],
        # 3. Недопустимый срок (не 1-50 лет и не 2000-2100 год) -> term_invalid.
        [
            UNIVERSITIES[1],
            VENDORS[2],
            PRODUCTS[2],
            "Д-902/2026",
            "да",
            "1500",
            "передано",
            MANAGERS[2],
            "",
            "",
        ],
    ]


def _license_rows() -> list[list[str]]:
    return [*_good_license_rows(27), *_bad_license_rows()]


def make_licenses_10_columns_xlsx() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Лицензии"
    sheet.append(LICENSE_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    for row in _license_rows():
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_licenses_10_columns_xls() -> bytes:
    workbook = xlwt.Workbook(encoding="utf-8")
    sheet = workbook.add_sheet("Лицензии")
    bold = xlwt.easyxf("font: bold on")
    for col, header in enumerate(LICENSE_HEADERS):
        sheet.write(0, col, header, bold)
    for row_idx, row in enumerate(_license_rows(), start=1):
        for col, value in enumerate(row):
            sheet.write(row_idx, col, value)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_licenses_multisheet() -> bytes:
    """Три листа по годам; на первом — вендор в объединённых ячейках (B3-2 `fill_down`)."""
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    bold = Font(bold=True)

    for year_idx, year in enumerate(("2024", "2025", "2026")):
        sheet = workbook.create_sheet(year)
        sheet.append([f"Каталог лицензий за {year} год"])
        sheet.append([])
        sheet.append(LICENSE_HEADERS)
        for cell in sheet[3]:
            cell.font = bold
        vendor = VENDORS[year_idx % len(VENDORS)]
        for i in range(3):
            university = UNIVERSITIES[(year_idx * 3 + i) % len(UNIVERSITIES)]
            sheet.append(
                [
                    university,
                    vendor if i == 0 else "",  # объединённая ячейка вендора на блок строк
                    PRODUCTS[i % len(PRODUCTS)],
                    f"Д-{200 + year_idx * 10 + i}/{year}",
                    "да",
                    "3",
                    "передано",
                    MANAGERS[i % len(MANAGERS)],
                    "",
                    "",
                ]
            )
        if year_idx == 0:
            # Объединяем колонку «Вендор» (B) для трёх строк данных (4..6) — реальный файл
            # заказчика часто пишет вендора один раз на блок, а не на каждую строку.
            sheet.merge_cells(start_row=4, start_column=2, end_row=6, end_column=2)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_directions() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Направления"
    sheet.append(["Код", "Название", "Родитель"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    rows = [
        ["qa", "Тестирование (QA)", ""],
        ["qa-auto", "Автоматизация тестирования", "qa"],
        ["be", "Backend-разработка", ""],
        ["be-python", "Backend на Python", "be"],
        ["data", "Данные и аналитика", ""],
    ]
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_responsible() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Ответственные"
    sheet.append(["Вуз", "ИНН вуза", "Почта менеджера", "ФИО менеджера"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    rows = [
        [UNIVERSITIES[0], "", "ivanov@example.com", ""],
        [UNIVERSITIES[1], "", "", "Петрова Анна Сергеевна"],
        [UNIVERSITIES[2], "", "sidorov@example.com", ""],
    ]
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


SAMPLES = {
    "licenses-10-columns.xlsx": make_licenses_10_columns_xlsx,
    "licenses-10-columns.xls": make_licenses_10_columns_xls,
    "licenses-multisheet.xlsx": make_licenses_multisheet,
    "directions.xlsx": make_directions,
    "responsible.xlsx": make_responsible,
}


def main() -> None:
    SAMPLES_DIR.mkdir(parents=True, exist_ok=True)
    for name, builder in SAMPLES.items():
        content = builder()
        path = SAMPLES_DIR / name
        path.write_bytes(content)
        print(f"{name}: {len(content)} bytes")


if __name__ == "__main__":
    main()
